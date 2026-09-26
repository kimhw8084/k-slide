from __future__ import annotations

import json
import struct
import tempfile
import unittest
import zlib
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from pathlib import Path

from PIL import Image

from evals.performance_load import _KNOWN_STDOUT_LIBRARY_WARNING, _parse_product_cli_json
from k_slide.cli import _charge_model_output_budget, _reserve_model_input_budget
from k_slide.certification import candidate_completeness
from k_slide.errors import ErrorCode, KSlideError
from k_slide.environment import RunEnvironmentIdentity
from k_slide.host_adapter import HostInputReference, validate_host_inputs
from k_slide.ingest import prepare_run
from k_slide.normalization import _preflight_resources, normalize_run
from k_slide.paas import (
    AuthorizedScopeContext,
    PaaSController,
    PaaSJobRequest,
    ReferencePaaSJobService,
    RuntimeIdentity,
    ScopedAdmissionPolicy,
)
from k_slide.resource_budget import (
    ResourceBudget,
    estimate_model_input_tokens,
)
from k_slide.state import RunPhase, load_state
from k_slide.storage import StorageArtifact, storage_path
from k_slide.io import read_json
from k_slide.security import validate_input
from k_slide.model import build_translation_prompt
from tests.reference_fixtures import reference_environment


def _budget(**overrides: int | None) -> ResourceBudget:
    values = dict(ResourceBudget.reference().values)
    values.update(overrides)
    return ResourceBudget(tuple(sorted(values.items())), "reference_non_production")


def _profile(root: Path, budget: ResourceBudget) -> None:
    path = root / ".k-slide-config" / "production-profile.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"schema_version": "1.1", "resource_budget": budget.as_dict()}), encoding="utf-8")


def _png(path: Path, width: int = 2, height: int = 2) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (width, height), "white").save(path, format="PNG")


def _png_header_only(path: Path, width: int, height: int) -> None:
    def chunk(tag: bytes, payload: bytes) -> bytes:
        return struct.pack(">I", len(payload)) + tag + payload + struct.pack(">I", zlib.crc32(tag + payload) & 0xFFFFFFFF)

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    raw = b"\x00\x00\x00\x00\x00"
    data = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")
    path.write_bytes(data)


def _runtime() -> RuntimeIdentity:
    environment = reference_environment("ksa36-paas-resource-budget")
    return RuntimeIdentity("ksa36-reference-runtime", "NOT_INVOKED", "native-only", "core-reference", environment_identity=environment)


class ResourceBudgetContractTests(unittest.TestCase):
    def test_load_harness_accepts_only_the_known_library_warning_before_cli_json(self) -> None:
        payload = json.dumps({"status": "EXTRACTED", "run_id": "synthetic-run"})
        self.assertEqual(_parse_product_cli_json(payload), json.loads(payload))
        self.assertEqual(_parse_product_cli_json(f"{_KNOWN_STDOUT_LIBRARY_WARNING}\n{payload}"), json.loads(payload))
        self.assertIsNone(_parse_product_cli_json(f"unexpected diagnostic\n{payload}"))
        self.assertIsNone(_parse_product_cli_json(f"{_KNOWN_STDOUT_LIBRARY_WARNING}\nnot-json"))

    def test_resource_budget_and_telemetry_error_schemas_match_runtime_contract(self) -> None:
        import jsonschema

        root = Path(__file__).resolve().parents[1]
        budget_schema = json.loads((root / "schemas" / "resource-budget.schema.json").read_text(encoding="utf-8"))
        telemetry_schema = json.loads((root / "schemas" / "operational-telemetry-event.schema.json").read_text(encoding="utf-8"))
        jsonschema.Draft202012Validator.check_schema(budget_schema)
        jsonschema.validate(ResourceBudget.reference().as_dict(), budget_schema)
        enum = telemetry_schema["properties"]["error_code"]["enum"]
        self.assertEqual(sorted(enum), sorted(code.value for code in ErrorCode))

    def test_every_numeric_limit_rejects_above_and_accepts_below_at(self) -> None:
        budget = ResourceBudget.reference()
        for name, limit in budget.values:
            if name in {"max_media_duration_seconds", "model_media_token_reserve_per_item"}:
                continue
            with self.subTest(resource=name):
                budget.enforce(name, max(0, int(limit) - 1))
                budget.enforce(name, int(limit))
                with self.assertRaises(KSlideError) as raised:
                    budget.enforce(name, int(limit) + 1)
                self.assertEqual(raised.exception.code, ErrorCode.RESOURCE_LIMIT)
        budget.enforce("max_media_duration_seconds", 0)
        with self.assertRaises(KSlideError) as duration:
            budget.enforce("max_media_duration_seconds", 1)
        self.assertEqual(duration.exception.code, ErrorCode.RESOURCE_LIMIT)

    def test_contract_is_versioned_strict_and_reference_defaults_are_nonproduction(self) -> None:
        budget = ResourceBudget.reference()
        self.assertEqual(ResourceBudget.from_dict(budget.as_dict()), budget)
        self.assertEqual(budget.environment, "reference_non_production")
        self.assertEqual(len(budget.sha256), 64)
        with self.assertRaises(KSlideError) as missing:
            ResourceBudget.from_dict({"schema_version": "1.0", "environment": "production", "limits": {}})
        self.assertEqual(missing.exception.code, ErrorCode.RESOURCE_BUDGET_INVALID)
        with self.assertRaises(KSlideError):
            ResourceBudget.from_dict({**budget.as_dict(), "schema_version": "9.0"})

    def test_production_budget_is_required_and_profile_reload_is_durable(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with self.assertRaises(KSlideError) as missing:
                ResourceBudget.load_for_workspace(root, production_required=True)
            self.assertEqual(missing.exception.code, ErrorCode.RESOURCE_BUDGET_UNAVAILABLE)
            _profile(root, ResourceBudget.reference())
            with self.assertRaises(KSlideError) as reference:
                ResourceBudget.load_for_workspace(root, production_required=True)
            self.assertEqual(reference.exception.code, ErrorCode.RESOURCE_BUDGET_INVALID)
            production = ResourceBudget(tuple(sorted(dict(ResourceBudget.reference().values).items())), "production")
            _profile(root, production)
            loaded = ResourceBudget.load_for_workspace(root, production_required=True)
            self.assertEqual(loaded, production)
            first_hash = loaded.sha256
            _profile(root, replace(production, values=tuple(sorted({**dict(production.values), "max_file_bytes": 1024}.items()))))
            self.assertNotEqual(ResourceBudget.load_for_workspace(root, production_required=True).sha256, first_hash)
            self.assertIn("resource_budget", candidate_completeness({}, "PRODUCTION_CERTIFIED"))
            with self.assertRaises(KSlideError) as legacy_managed_run:
                ResourceBudget.from_run_manifest(root, {}, production_required=True)
            self.assertEqual(legacy_managed_run.exception.code, ErrorCode.RESOURCE_BUDGET_UNAVAILABLE)

    def test_intake_file_and_host_adapters_share_the_same_budget_result(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            source = workspace / "sample.png"
            _png(source)
            byte_count = source.stat().st_size
            outcomes: list[str | None] = []
            for topology in ("local_cli", "opencode", "cloud_vscode"):
                root = workspace / topology
                root.mkdir()
                local_source = root / "sample.png"
                local_source.write_bytes(source.read_bytes())
                _profile(root, _budget(max_file_bytes=byte_count - 1))
                if topology == "local_cli":
                    run = prepare_run(root, explicit_paths=[str(local_source)], environment_identity=reference_environment("ksa36-intake-parity"))
                else:
                    kind = "attachment" if topology == "opencode" else "workspace_file"
                    run = prepare_run(root, host_input_refs=[HostInputReference(kind, local_source.name, str(local_source))], approved_root=root, environment_identity=reference_environment("ksa36-intake-parity"))
                state = load_state(run)
                outcomes.append(state.error_code)
                self.assertEqual(state.phase, RunPhase.FAILED_INPUT)
            self.assertEqual(outcomes, [ErrorCode.INPUT_TOO_LARGE.value] * 3)

            for max_bytes in (byte_count, byte_count + 1):
                root = workspace / f"pass-{max_bytes}"
                root.mkdir()
                accepted = root / "sample.png"
                accepted.write_bytes(source.read_bytes())
                _profile(root, _budget(max_file_bytes=max_bytes))
                run = prepare_run(root, explicit_paths=[str(accepted)], environment_identity=reference_environment("ksa36-boundary"))
                self.assertEqual(load_state(run).phase, RunPhase.INPUT_VALIDATED)
                manifest = read_json(storage_path(run, StorageArtifact.RUN_MANIFEST, "RUN_MANIFEST.json"))
                self.assertEqual(manifest["resource_budget_sha256"], _budget(max_file_bytes=max_bytes).sha256)

    def test_total_input_bytes_and_document_count_fail_before_snapshot(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source_dir = root / "inputs"
            first, second = source_dir / "one.png", source_dir / "two.png"
            _png(first)
            _png(second)
            total = first.stat().st_size + second.stat().st_size
            _profile(root, _budget(max_total_file_bytes_per_run=total - 1))
            run = prepare_run(root, explicit_paths=[str(first), str(second)], environment_identity=reference_environment("ksa36-total-bytes"))
            self.assertEqual(load_state(run).error_code, ErrorCode.RESOURCE_LIMIT.value)
            self.assertFalse((run / "inputs" / f"source-001{first.suffix}").exists())
            _profile(root, _budget(max_documents_per_run=1))
            run = prepare_run(root, explicit_paths=[str(first), str(second)], environment_identity=reference_environment("ksa36-documents"))
            self.assertEqual(load_state(run).error_code, ErrorCode.RESOURCE_LIMIT.value)
            self.assertFalse((run / "inputs" / f"source-001{first.suffix}").exists())

    def test_run_uses_immutable_budget_snapshot_after_profile_reload(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "sample.png"
            _png(source, 3, 3)
            original = _budget(max_pixels_per_image=100)
            _profile(root, original)
            environment = reference_environment("ksa36-durable-budget")
            run = prepare_run(root, explicit_paths=[str(source)], environment_identity=environment)
            _profile(root, _budget(max_pixels_per_image=1))
            result = normalize_run(run, environment_identity=environment)
            self.assertEqual(sum(unit.width_px * unit.height_px for document in result.documents for unit in document.units), 9)

    def test_normalized_byte_budget_boundary_is_rechecked_after_render(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "sample.png"
            _png(source, 3, 3)
            env = reference_environment("ksa36-normalized-bytes")
            baseline = prepare_run(root, explicit_paths=[str(source)], environment_identity=env)
            baseline_result = normalize_run(baseline, environment_identity=env)
            output = baseline / baseline_result.documents[0].units[0].canonical_render_path
            normalized_bytes = output.stat().st_size
            self.assertGreater(normalized_bytes, 0)

            for limit, expected in ((normalized_bytes - 1, False), (normalized_bytes, True), (normalized_bytes + 1, True)):
                case = root / f"case-{limit}"
                case.mkdir()
                case_source = case / "sample.png"
                case_source.write_bytes(source.read_bytes())
                _profile(case, _budget(max_normalized_bytes_per_run=limit))
                run = prepare_run(case, explicit_paths=[str(case_source)], environment_identity=env)
                if expected:
                    result = normalize_run(run, environment_identity=env)
                    produced = run / result.documents[0].units[0].canonical_render_path
                    self.assertLessEqual(produced.stat().st_size, limit)
                else:
                    with self.assertRaises(KSlideError) as raised:
                        normalize_run(run, environment_identity=env)
                    self.assertEqual(raised.exception.code, ErrorCode.RESOURCE_LIMIT)

    def test_image_pixel_and_decompression_bomb_preflight(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            below = root / "below.png"
            at = root / "at.png"
            above = root / "above.png"
            _png(below, 99, 100)
            _png(at, 100, 100)
            _png(above, 101, 100)
            budget = _budget(max_pixels_per_image=10_000, max_total_decoded_pixels_per_run=100_000)
            _preflight_resources([("source-001", below, {"extension": ".png"})], budget)
            _preflight_resources([("source-001", at, {"extension": ".png"})], budget)
            with self.assertRaises(KSlideError) as raised:
                _preflight_resources([("source-001", above, {"extension": ".png"})], budget)
            self.assertEqual(raised.exception.code, ErrorCode.INPUT_TOO_LARGE)

            dimension_limited = root / "dimension-limited.png"
            _png(dimension_limited, 101, 2)
            with self.assertRaises(KSlideError) as dimension:
                _preflight_resources([("source-001", dimension_limited, {"extension": ".png"})], _budget(max_image_dimension=100, max_pixels_per_image=10_000))
            self.assertEqual(dimension.exception.code, ErrorCode.INPUT_TOO_LARGE)

            bomb = root / "bomb.png"
            _png_header_only(bomb, 30_000, 30_000)
            with self.assertRaises(KSlideError) as compressed:
                _preflight_resources([("source-001", bomb, {"extension": ".png"})], ResourceBudget.reference())
            self.assertEqual(compressed.exception.code, ErrorCode.RESOURCE_LIMIT)

    def test_animated_webp_is_rejected_without_silent_frame_truncation(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "animated.webp"
            frames = [Image.new("RGB", (4, 4), color) for color in ("red", "blue")]
            frames[0].save(source, format="WEBP", save_all=True, append_images=frames[1:], duration=50, loop=0)
            with self.assertRaises(KSlideError) as rejected:
                _preflight_resources([("source-001", source, {"extension": ".webp"})], ResourceBudget.reference())
            self.assertEqual(rejected.exception.code, ErrorCode.INPUT_UNSUPPORTED)

    def test_pdf_page_slide_unit_pixel_and_normalized_byte_boundaries(self) -> None:
        try:
            import fitz
        except ImportError:
            self.skipTest("PyMuPDF is unavailable")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf = root / "deck.pdf"
            document = fitz.open()
            for _ in range(3):
                document.new_page(width=100, height=100)
            document.save(pdf)
            document.close()
            snapshots = [("source-001", pdf, {"extension": ".pdf"})]
            _preflight_resources(snapshots, _budget(max_pdf_pages_per_document=4, max_work_units_per_run=4, max_total_decoded_pixels_per_run=1_000_000))
            _preflight_resources(snapshots, _budget(max_pdf_pages_per_document=3, max_work_units_per_run=3, max_total_decoded_pixels_per_run=1_000_000))
            for field in ("max_pdf_pages_per_document", "max_work_units_per_run"):
                with self.subTest(resource=field), self.assertRaises(KSlideError) as raised:
                    _preflight_resources(snapshots, _budget(**{field: 2, "max_total_decoded_pixels_per_run": 1_000_000}))
                self.assertEqual(raised.exception.code, ErrorCode.RESOURCE_LIMIT)
            with self.assertRaises(KSlideError) as pixel_limit:
                _preflight_resources(snapshots, _budget(max_total_decoded_pixels_per_run=1))
            self.assertEqual(pixel_limit.exception.code, ErrorCode.RESOURCE_LIMIT)
            pixels = 3 * 306 * 306
            _preflight_resources(snapshots, _budget(max_total_decoded_pixels_per_run=pixels))
            with self.assertRaises(KSlideError):
                _preflight_resources(snapshots, _budget(max_total_decoded_pixels_per_run=pixels - 1))
            with self.assertRaises(KSlideError) as per_page:
                _preflight_resources(snapshots, _budget(max_image_dimension=305, max_total_decoded_pixels_per_run=pixels))
            self.assertEqual(per_page.exception.code, ErrorCode.RESOURCE_LIMIT)

    def test_pptx_slide_count_is_checked_before_office_conversion(self) -> None:
        try:
            from pptx import Presentation
        except ImportError:
            self.skipTest("python-pptx is unavailable")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pptx_path = root / "deck.pptx"
            presentation = Presentation()
            for _ in range(3):
                presentation.slides.add_slide(presentation.slide_layouts[6])
            presentation.save(pptx_path)
            snapshots = [("source-001", pptx_path, {"extension": ".pptx"})]
            _preflight_resources(snapshots, _budget(max_pptx_slides_per_deck=4, max_work_units_per_run=4, max_total_decoded_pixels_per_run=100_000_000))
            _preflight_resources(snapshots, _budget(max_pptx_slides_per_deck=3, max_work_units_per_run=3, max_total_decoded_pixels_per_run=100_000_000))
            with self.assertRaises(KSlideError) as raised:
                _preflight_resources(snapshots, _budget(max_pptx_slides_per_deck=2, max_work_units_per_run=3, max_total_decoded_pixels_per_run=100_000_000))
            self.assertEqual(raised.exception.code, ErrorCode.RESOURCE_LIMIT)

    def test_model_media_and_token_budgets_block_without_content_telemetry(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            run = Path(directory) / "run"
            run.mkdir()
            metrics_path = storage_path(run, StorageArtifact.METRICS, "metrics.json", create_parent=True)
            packet = {"work_unit_id": "doc-001-page-0001", "source_regions": [{"english": "synthetic"}]}
            base = ResourceBudget.reference()
            exact_input = estimate_model_input_tokens(packet, build_translation_prompt(), 1, base)
            allowed = _budget(
                max_media_items_per_work_unit=1,
                max_model_input_tokens_per_work_unit=exact_input,
                max_model_tokens_per_run=exact_input + 200,
                max_model_output_tokens_per_work_unit=200,
            )
            reserved = _reserve_model_input_budget(run, work_unit_id="doc-001-page-0001", evidence_revision="a" * 64, packet=packet, media_count=1, budget=allowed)
            self.assertEqual(reserved["input_tokens_estimated"], exact_input)
            with self.assertRaises(KSlideError) as input_limit:
                _reserve_model_input_budget(
                    Path(directory) / "input-limit",
                    work_unit_id="doc-003-page-0001",
                    evidence_revision="c" * 64,
                    packet=packet,
                    media_count=1,
                    budget=_budget(max_model_input_tokens_per_work_unit=exact_input - 1, max_model_tokens_per_run=exact_input + 100, max_model_output_tokens_per_work_unit=100),
                )
            self.assertEqual(input_limit.exception.code, ErrorCode.RESOURCE_LIMIT)
            with self.assertRaises(KSlideError) as media:
                _reserve_model_input_budget(run, work_unit_id="doc-002-page-0001", evidence_revision="b" * 64, packet=packet, media_count=2, budget=allowed)
            self.assertEqual(media.exception.code, ErrorCode.RESOURCE_LIMIT)

            marker = "SYNTHETIC_TRANSLATION_SECRET"
            output = json.dumps({"english": marker}, separators=(",", ":"))
            exact_output_budget = _budget(max_model_output_tokens_per_work_unit=len(output), max_model_tokens_per_run=exact_input + len(output))
            exact_output_run = Path(directory) / "exact-output"
            exact_output_run.mkdir()
            self.assertEqual(_charge_model_output_budget(exact_output_run, output, exact_output_budget), len(output))
            _charge_model_output_budget(run, output, allowed)
            metrics_text = metrics_path.read_text(encoding="utf-8")
            self.assertNotIn(marker, metrics_text)
            output_budget = _budget(max_model_output_tokens_per_work_unit=len(output) - 1)
            rejected_run = Path(directory) / "rejected"
            rejected_run.mkdir()
            with self.assertRaises(KSlideError) as oversized:
                _charge_model_output_budget(rejected_run, output, output_budget)
            self.assertEqual(oversized.exception.code, ErrorCode.RESOURCE_LIMIT)
            self.assertNotIn(marker, storage_path(rejected_run, StorageArtifact.METRICS, "metrics.json").read_text(encoding="utf-8"))

    def test_shared_scoped_queue_capacity_is_race_safe_and_survives_reload(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            budget = _budget(max_queued_runs_per_scope=2)
            policy = ScopedAdmissionPolicy.from_resource_budget(budget)
            scope = AuthorizedScopeContext("ksa36-user", "ksa36-workspace", "ksa36-scope")
            runtime = _runtime()
            service = ReferencePaaSJobService(root, admission_policy=policy, scope_context=scope)
            controller = PaaSController(service, scope_context=scope)

            def submit(index: int) -> str:
                request = PaaSJobRequest(
                    f"ksa36-queue-{index}",
                    str(scope.scope_ref),
                    f"ksa36-store-{index}",
                    runtime,
                    total_work_units=3,
                    scope_context=scope,
                )
                try:
                    return controller.submit(request).status.value
                except KSlideError as exc:
                    return exc.code.value

            with ThreadPoolExecutor(max_workers=5) as pool:
                outcomes = list(pool.map(submit, range(5)))
            self.assertEqual(outcomes.count(ErrorCode.CAPACITY_LIMIT.value), 2)
            queue = service.queue(scope_context=scope)
            self.assertEqual(sum(entry.active for entry in queue.entries), 1)
            self.assertEqual(len(queue.queued), 2)

            reloaded_budget = ResourceBudget.from_dict(budget.as_dict())
            reloaded = ReferencePaaSJobService(root, admission_policy=ScopedAdmissionPolicy.from_resource_budget(reloaded_budget), scope_context=scope)
            self.assertEqual(len(reloaded.queue(scope_context=scope).queued), 2)
            next_request = PaaSJobRequest("ksa36-queue-after-reload", str(scope.scope_ref), "ksa36-store-after-reload", runtime, total_work_units=1, scope_context=scope)
            with self.assertRaises(KSlideError) as full:
                PaaSController(reloaded, scope_context=scope).submit(next_request)
            self.assertEqual(full.exception.code, ErrorCode.CAPACITY_LIMIT)

    def test_shared_host_entrypoints_fail_closed_on_local_run_capacity_race(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            source = workspace / "sample.png"
            _png(source)
            _profile(workspace, _budget(max_concurrent_runs_per_scope=1))
            env = reference_environment("ksa36-local-capacity")

            def prepare(index: int) -> tuple[RunPhase, str | None]:
                if index % 2:
                    run = prepare_run(
                        workspace,
                        host_input_refs=[HostInputReference("attachment", source.name, str(source))],
                        approved_root=workspace,
                        environment_identity=env,
                    )
                else:
                    run = prepare_run(workspace, explicit_paths=[str(source)], environment_identity=env)
                state = load_state(run)
                return state.phase, state.error_code

            with ThreadPoolExecutor(max_workers=2) as pool:
                outcomes = list(pool.map(prepare, range(2)))
            admitted = sum(phase is RunPhase.INPUT_VALIDATED for phase, _ in outcomes)
            capacity_limited = sum(code == ErrorCode.CAPACITY_LIMIT.value for _, code in outcomes)
            self.assertEqual(admitted, 1)
            self.assertEqual(capacity_limited, 1)

    def test_host_reference_rejection_carries_typed_budget_error(self) -> None:
        reference = HostInputReference("attachment", "blocked.png", "/does/not/exist.png", rejection_code=ErrorCode.RESOURCE_LIMIT.value)
        with self.assertRaises(KSlideError) as rejected:
            validate_host_inputs([reference], approved_root=Path.cwd())
        self.assertEqual(rejected.exception.code, ErrorCode.RESOURCE_LIMIT)


if __name__ == "__main__":
    unittest.main()
