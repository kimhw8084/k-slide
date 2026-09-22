from __future__ import annotations

import json
import tempfile
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path
from types import SimpleNamespace

from k_slide.cli import _conflict_assess, _next, _submit
from k_slide.errors import ErrorCode, KSlideError
from k_slide.evidence_ir import EvidenceIR, EvidenceRegion
from k_slide.extraction import _tables
from k_slide.ingest import prepare_run
from k_slide.normalization import _connector_endpoints, _pptx_native
from k_slide.translation import parse_translation_patch
from k_slide.verify import finalize_run, verify_run
from tests.reference_fixtures import reference_environment


FIXTURES = Path(__file__).parent / "fixtures"


def _evidence(visual_elements: tuple[dict, ...], *, source_text: str = "Visual", language: str | None = None) -> EvidenceIR:
    return EvidenceIR(
        "doc",
        "unit",
        {},
        regions=(EvidenceRegion("unit-r", selected_literal_candidate=source_text, language=language),),
        visual_elements=visual_elements,
        required_source_ids=("unit-r", *(item["element_id"] for item in visual_elements)),
    ).with_revision()


def _patch(evidence: EvidenceIR, *, relation: dict | None = None, english: str = "Visual") -> dict:
    return {
        "schema_version": "1.0",
        "work_unit_id": evidence.work_unit_id,
        "evidence_revision": evidence.evidence_revision,
        "regions": [{"region_id": "unit-r", "english": english, "term_ids": [], "unresolved": False}],
        "tables": [],
        "visual_interpretations": [relation] if relation else [],
        "executive_claims": [],
    }


class KSA24F24RepairTests(unittest.TestCase):
    def test_connector_direction_requires_explicit_arrow_evidence_and_supports_reverse_start_arrow(self) -> None:
        fixture = json.loads((FIXTURES / "ksa24_connector_fixture.json").read_text(encoding="utf-8"))
        elements = tuple(fixture["elements"] + fixture["connectors"])
        evidence = _evidence(elements)

        def relation(edge: str, endpoints: tuple[str, str], direction: str) -> dict:
            return {
                "relation_id": edge,
                "interpretation": "flow",
                "evidence_ids": [edge],
                "source_element_ids": [*endpoints, edge],
                "relation_type": "next",
                "direction": direction,
                "provenance": "source_fact",
            }

        with self.assertRaises(KSlideError) as no_arrow:
            parse_translation_patch(_patch(evidence, relation=relation("fixture-no-arrow", ("fixture-group-a", "fixture-group-b"), "left_to_right"))).validate_against(evidence)
        self.assertEqual(no_arrow.exception.code, ErrorCode.SCHEMA_INVALID)

        end_arrow = relation("fixture-end-arrow", ("fixture-a", "fixture-b"), "left_to_right")
        parse_translation_patch(_patch(evidence, relation=end_arrow)).validate_against(evidence)

        reverse = relation("fixture-start-arrow", ("fixture-b", "fixture-a"), "right_to_left")
        parse_translation_patch(_patch(evidence, relation=reverse)).validate_against(evidence)
        with self.assertRaises(KSlideError):
            parse_translation_patch(_patch(evidence, relation=relation("fixture-start-arrow", ("fixture-a", "fixture-b"), "left_to_right"))).validate_against(evidence)

        with self.assertRaises(KSlideError):
            parse_translation_patch(_patch(evidence, relation=relation("fixture-bidirectional", ("fixture-a", "fixture-b"), "left_to_right"))).validate_against(evidence)
        parse_translation_patch(_patch(evidence, relation=relation("fixture-bidirectional", ("fixture-a", "fixture-b"), "bidirectional"))).validate_against(evidence)
        with self.assertRaises(KSlideError):
            parse_translation_patch(_patch(evidence, relation=relation("fixture-ambiguous", ("fixture-a", "fixture-b"), "left_to_right"))).validate_against(evidence)

    def test_normalizer_reads_head_and_tail_arrowheads_from_ooxml(self) -> None:
        xml = ET.fromstring(
            '<p:cxnSp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<p:nvCxnSpPr><p:cNvCxnSpPr><a:stCxn id="1"/><a:endCxn id="2"/></p:cNvCxnSpPr></p:nvCxnSpPr>'
            '<p:spPr><a:ln><a:headEnd type="triangle"/><a:tailEnd type="none"/></a:ln></p:spPr></p:cxnSp>'
        )
        shape = SimpleNamespace(_element=xml)
        result = _connector_endpoints(shape, {1: "a", 2: "b"})
        self.assertIsNotNone(result)
        assert result is not None
        self.assertEqual(result["direction_evidence"], "end_to_start")
        self.assertEqual(result["start_arrow_type"], "triangle")
        self.assertEqual(result["end_arrow_type"], "none")

    def test_chart_claims_are_closed_to_series_point_category_and_recomputed_fact(self) -> None:
        chart = json.loads((FIXTURES / "ksa24_chart_claim_fixture.json").read_text(encoding="utf-8"))
        evidence = _evidence((chart,))

        def claim(**changes: object) -> dict:
            value = {
                "chart_element_id": "fixture-chart",
                "kind": "trend",
                "series_index": 0,
                "series_name": "North",
                "direction": "increasing",
            }
            value.update(changes)
            return value

        good = {"relation_id": "trend", "interpretation": "North is increasing", "evidence_ids": ["fixture-chart"], "source_element_ids": ["fixture-chart"], "provenance": "source_fact", "chart_claim": claim()}
        parse_translation_patch(_patch(evidence, relation=good)).validate_against(evidence)
        for bad in (
            claim(series_index=1),
            claim(direction="decreasing"),
            claim(series_name="South"),
        ):
            with self.assertRaises(KSlideError):
                parse_translation_patch(_patch(evidence, relation={**good, "chart_claim": bad})).validate_against(evidence)

        value_claim = claim(kind="value", point_index=0, category="Q1", value=-1234.5)
        parse_translation_patch(_patch(evidence, relation={**good, "chart_claim": value_claim}, english="Visual")).validate_against(evidence)
        for bad in (
            {**value_claim, "value": -1234.4},
            {**value_claim, "category": "Q2"},
        ):
            with self.assertRaises(KSlideError):
                parse_translation_patch(_patch(evidence, relation={**good, "chart_claim": bad})).validate_against(evidence)

        flat = claim(series_index=2, series_name="Flat", direction="flat")
        parse_translation_patch(_patch(evidence, relation={**good, "chart_claim": flat})).validate_against(evidence)
        nullable = claim(series_index=3, series_name="Nullable", direction="increasing")
        with self.assertRaises(KSlideError):
            parse_translation_patch(_patch(evidence, relation={**good, "chart_claim": nullable})).validate_against(evidence)
        ranking = claim(kind="ranking", series_index=0, series_name="North", point_index=0, category="Q1", ranking="lowest", rank=1)
        parse_translation_patch(_patch(evidence, relation={**good, "chart_claim": ranking})).validate_against(evidence)

        with self.assertRaises(KSlideError):
            parse_translation_patch(_patch(evidence, relation={**good, "provenance": "source_fact", "chart_claim": None})).validate_against(evidence)
        with self.assertRaises(KSlideError):
            parse_translation_patch(_patch(evidence, relation={"relation_id": "text", "interpretation": "North is decreasing while South is increasing", "evidence_ids": ["fixture-chart"], "source_element_ids": ["fixture-chart"], "provenance": "supported_interpretation"})).validate_against(evidence)

        numeric_text = {"relation_id": "number", "interpretation": "The value is -1,234.5%", "evidence_ids": ["fixture-chart"], "source_element_ids": ["fixture-chart"], "provenance": "supported_interpretation"}
        parse_translation_patch(_patch(evidence, relation=numeric_text)).validate_against(evidence)
        with self.assertRaises(KSlideError):
            parse_translation_patch(_patch(evidence, relation={**numeric_text, "interpretation": "The value is -1,234.4%"})).validate_against(evidence)

    def test_modality_matrix_requires_status_and_rejects_english_strengthening(self) -> None:
        cases = (
            ("결정", "decided", "decision", "The decision is approved."),
            ("약속", "committed", "status", "The work is committed."),
            ("계획", "planned", "status", "The work is planned."),
            ("예정", "scheduled", "status", "The work is scheduled."),
            ("목표", "target", "status", "The target is the target."),
            ("제안", "proposed", "status", "The proposal is proposed."),
            ("검토 중", "under_review", "status", "The work is under review."),
            ("가능", "possible", "status", "The work is possible."),
            ("잠정", "tentative", "status", "The result is tentative."),
            ("전망", "forecast", "status", "The forecast is projected."),
            ("완료", "completed", "status", "The work is completed."),
            ("진행 중", "in_progress", "status", "The work is in progress."),
        )
        for source_text, status, speech, english in cases:
            with self.subTest(source_text=source_text):
                evidence = _evidence((), source_text=source_text, language="ko")
                region = {"region_id": "unit-r", "english": english, "term_ids": [], "unresolved": False, "commitment_status": status, "speech_act": speech, "evidence_ids": ["unit-r"]}
                parse_translation_patch({**_patch(evidence), "regions": [region]}).validate_against(evidence)
                omitted = dict(region)
                omitted.pop("commitment_status")
                with self.assertRaises(KSlideError):
                    parse_translation_patch({**_patch(evidence), "regions": [omitted]}).validate_against(evidence)

        decided = _evidence((), source_text="결정", language="ko")
        missing_speech = {"region_id": "unit-r", "english": "The decision is approved.", "term_ids": [], "unresolved": False, "commitment_status": "decided", "evidence_ids": ["unit-r"]}
        with self.assertRaises(KSlideError):
            parse_translation_patch({**_patch(decided), "regions": [missing_speech]}).validate_against(decided)
        under_review = _evidence((), source_text="검토 중", language="ko")
        strengthened = {"region_id": "unit-r", "english": "The work is approved.", "term_ids": [], "unresolved": False, "commitment_status": "under_review", "speech_act": "status", "evidence_ids": ["unit-r"]}
        with self.assertRaises(KSlideError):
            parse_translation_patch({**_patch(under_review), "regions": [strengthened]}).validate_against(under_review)
        for source_text, wrong_status in (("계획", "committed"), ("예정", "decided"), ("목표", "decided"), ("약속", "tentative"), ("약속", "decided")):
            evidence = _evidence((), source_text=source_text, language="ko")
            wrong = {"region_id": "unit-r", "english": "The work is stated.", "term_ids": [], "unresolved": False, "commitment_status": wrong_status, "speech_act": "status", "evidence_ids": ["unit-r"]}
            with self.assertRaises(KSlideError):
                parse_translation_patch({**_patch(evidence), "regions": [wrong]}).validate_against(evidence)

    def test_table_unit_derivation_blank_and_ambiguous_scope(self) -> None:
        cells = [
            {"cell_id": "t-r0-c0", "row": 0, "column": 0, "text": "KRW 억", "is_header": True},
            {"cell_id": "t-r0-c1", "row": 0, "column": 1, "text": "   ", "is_header": True},
            {"cell_id": "t-r1-c0", "row": 1, "column": 0, "text": "5", "is_header": False},
            {"cell_id": "t-r1-c1", "row": 1, "column": 1, "text": "\t", "is_header": False},
            {"cell_id": "t-r2-c0", "row": 2, "column": 0, "text": None, "is_spanned": True, "is_merge_origin": False},
        ]
        tables, facts = _tables(SimpleNamespace(work_unit_id="unit"), [{"source_id": "shape", "table": {"row_count": 3, "column_count": 2, "header_rows": [0], "cells": cells}}])
        self.assertEqual(tables[0].unit, "KRW 억")
        self.assertEqual(tables[0].cells[1].cell_state, "blank")
        self.assertEqual(tables[0].cells[4].cell_state, "merge_continuation")
        self.assertEqual(facts[0]["source_unit"], "KRW 억")
        self.assertEqual(facts[0]["canonical_value"], 5 * 10**8)
        self.assertEqual(facts[0]["currency"], "KRW")

        ambiguous = [dict(cells[0], text="USD millions"), dict(cells[1]), dict(cells[2]), dict(cells[3]), dict(cells[4])]
        ambiguous.insert(1, {"cell_id": "t-r0-c2", "row": 0, "column": 2, "text": "KRW 억", "is_header": True})
        tables, facts = _tables(SimpleNamespace(work_unit_id="unit"), [{"source_id": "shape", "table": {"row_count": 3, "column_count": 3, "header_rows": [0], "cells": ambiguous}}])
        self.assertIsNone(tables[0].unit)
        self.assertIsNone(facts[0]["source_unit"])

    def test_real_pptx_table_normalization_emits_header_unit_and_whitespace_blank(self) -> None:
        try:
            from PIL import Image
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError:
            self.skipTest("PPTX normalization dependencies are unavailable")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            shape = slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(5), Inches(2))
            table = shape.table
            table._tbl.tblPr.set("firstRow", "1")
            table.cell(0, 0).text = "USD millions"
            table.cell(0, 1).text = "Metric"
            table.cell(1, 0).text = "1.5"
            table.cell(1, 1).text = "   "
            table.cell(2, 0).text = "Revenue"
            table.cell(2, 1).text = "2"
            source = root / "fixture.pptx"
            presentation.save(source)
            render = root / "run" / "normalized" / "slide.png"
            render.parent.mkdir(parents=True)
            Image.new("RGB", (1280, 720), "white").save(render)
            units = _pptx_native(source, root / "run", "source-001", "doc-001", [render])
            native_table = next(item["table"] for item in units[0].native_evidence if "table" in item)
            self.assertEqual(native_table["unit"], "USD millions")
            whitespace = next(cell for cell in native_table["cells"] if cell["row"] == 1 and cell["column"] == 1)
            self.assertTrue(whitespace["is_blank"])
            tables, _facts = _tables(units[0], units[0].native_evidence)
            self.assertEqual(tables[0].cells[3].cell_state, "blank")

    def test_typed_runtime_submission_conflict_assessment_verify_and_finalize(self) -> None:
        try:
            from PIL import Image
        except ImportError:
            self.skipTest("Pillow is unavailable")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "slide.png"
            Image.new("RGB", (40, 40), "white").save(source)
            environment = reference_environment("ksa24-f24")
            run = prepare_run(root, explicit_paths=[str(source)], perform_processing=True, environment_identity=environment)
            next_value = _next(root, run.name, None, environment)
            self.assertEqual(next_value["status"], "READY")
            evidence = json.loads((run / "evidence" / f"{next_value['work_unit_id']}.json").read_text(encoding="utf-8"))
            payload = {
                "schema_version": "1.0",
                "work_unit_id": next_value["work_unit_id"],
                "evidence_revision": evidence["evidence_revision"],
                "regions": [{"region_id": evidence["regions"][0]["region_id"], "english": "", "term_ids": [], "unresolved": False}],
                "tables": [],
                "visual_interpretations": [],
                "executive_claims": [],
            }
            self.assertEqual(_submit(root, run.name, json.dumps(payload), None, environment)["status"], "ACCEPTED")
            self.assertEqual(_conflict_assess(root, run.name, '{"schema_version":"1.0","candidate_groups":[]}', None, environment)["status"], "ASSESSED_ZERO_CONFLICTS")
            verified = verify_run(run, environment_identity=environment)
            self.assertTrue(verified.passed, verified.as_dict())
            completed = finalize_run(run, environment_identity=environment)
            self.assertTrue(completed.passed)
            self.assertTrue((run / "RUN_COMPLETE.md").is_file())


if __name__ == "__main__":
    unittest.main()
