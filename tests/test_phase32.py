from __future__ import annotations

import json
import os
import stat
import tempfile
import unittest
from pathlib import Path

from evals.deck_scenarios import deck_scenarios
from evals.experiments import compare_aggregate
from evals.model_scorers import score_translation_patch
from evals.opencode_events import forbidden_tool_attempts, media_compliance, normalize_events
from evals.scenarios import PROTECTED_CATEGORIES, scenario_specs, split_manifest
from k_slide.evidence_ir import EvidenceIR, EvidenceRegion, EvidenceTable, EvidenceTableCell


def _evidence(*, regions=(), tables=(), numeric_facts=(), required_source_ids=()):
    return EvidenceIR(
        "doc-001",
        "doc-001-slide-0001",
        {"width_px": 1000, "height_px": 600},
        tuple(regions),
        tuple(tables),
        tuple(numeric_facts),
        required_source_ids=tuple(required_source_ids),
    ).with_revision()


class Phase32HarnessTests(unittest.TestCase):
    def test_stratified_splits_are_deterministic_and_protected(self) -> None:
        first = split_manifest()
        second = split_manifest()
        self.assertEqual(first, second)
        self.assertEqual(first["splits"], {"development": 60, "validation": 20, "held_out": 20})
        self.assertEqual(len(first["scenarios"]), 100)
        self.assertEqual(len({item["scenario_id"] for item in first["scenarios"]}), 100)
        for category in PROTECTED_CATEGORIES:
            self.assertGreater(first["by_category"][category]["held_out"], 0)
            self.assertEqual(sum(first["by_category"][category].values()), sum(1 for scenario in scenario_specs() if scenario.category == category))

    def test_split_manifest_is_emitted_with_scenario_membership(self) -> None:
        from evals.scenarios import write_specs

        with tempfile.TemporaryDirectory() as directory:
            write_specs(Path(directory))
            manifest = json.loads((Path(directory) / "splits.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["dataset_version"], "1.0")
            self.assertEqual(len(manifest["scenarios"]), 100)
            self.assertEqual({item["split"] for item in manifest["scenarios"]}, {"development", "validation", "held_out"})

    def test_prompt_text_is_not_a_forbidden_tool_attempt(self) -> None:
        events = normalize_events([{"type": "text", "text": "Do not use bash, write, or websearch."}])
        self.assertEqual(forbidden_tool_attempts(events), ())

    def test_actual_forbidden_tool_event_is_detected(self) -> None:
        events = normalize_events([{"type": "tool_use", "part": {"type": "tool", "tool": "bash", "callID": "c1", "state": {"input": {"cmd": "id"}}}}])
        self.assertEqual(forbidden_tool_attempts(events), ("bash",))

    def test_media_path_text_is_not_a_media_read(self) -> None:
        events = normalize_events([{"type": "text", "text": "Read .k-slide-runs/run/regions/context.png"}])
        compliance = media_compliance(events)
        self.assertFalse(compliance["required_context_image_read"])
        self.assertFalse(compliance["planned"])

    def test_external_read_is_not_inside_runtime_allowlist(self) -> None:
        from evals.opencode_runner import _read_policy_violations

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            events = normalize_events([{"type": "tool_use", "part": {"type": "tool", "tool": "read", "state": {"input": {"filePath": "/etc/passwd"}}}}])
            self.assertEqual(_read_policy_violations(root, events), ["/etc/passwd"])

    def test_actual_reads_are_required_before_submit(self) -> None:
        plan = {"model_media_plan": {"context_image": {"path": ".k-slide-runs/run/context.png", "required": True}, "required_crops": [{"region_id": "r1", "path": ".k-slide-runs/run/r1.png"}]}}
        events = normalize_events([
            {"type": "tool_result", "tool": "kslide_evidence", "output": json.dumps(plan)},
            {"type": "tool_use", "part": {"type": "tool", "tool": "read", "callID": "r1", "state": {"input": {"filePath": "/tmp/work/.k-slide-runs/run/context.png"}}}},
            {"type": "tool_use", "part": {"type": "tool", "tool": "read", "callID": "r2", "state": {"input": {"filePath": "/tmp/work/.k-slide-runs/run/r1.png"}}}},
            {"type": "tool_use", "part": {"type": "tool", "tool": "kslide_submit", "callID": "s1", "state": {"input": {}}}},
        ])
        compliance = media_compliance(events)
        self.assertEqual(compliance["required_crop_recall"], 1.0)
        self.assertTrue(compliance["required_context_image_read"])
        self.assertTrue(compliance["media_sequence_valid"])

    def test_success_contract_requires_complete_run(self) -> None:
        # The structured contract is also tested through the runner's helper
        # shape: an exit code alone cannot establish certification.
        from evals.opencode_runner import _completion_contract

        with tempfile.TemporaryDirectory() as directory:
            run = Path(directory)
            (run / "RUN_STATE.json").write_text(json.dumps({"phase": "VERIFIED"}), encoding="utf-8")
            self.assertFalse(_completion_contract(run)[0])
            (run / "RUN_STATE.json").write_text(json.dumps({"phase": "COMPLETE"}), encoding="utf-8")
            for name in ("RUN_COMPLETE.md", "05_executive_brief.md", "05_final_report.md", "06_verification.md", "07_unresolved_items.md"):
                (run / name).write_text("ok", encoding="utf-8")
            self.assertTrue(_completion_contract(run)[0])

    def test_opencode_exit_zero_without_k_slide_complete_fails(self) -> None:
        from evals.opencode_runner import OpenCodeEvalRunner

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            fake = root / "fake-opencode"
            fake.write_text("#!/bin/sh\nprintf '{\"type\":\"text\",\"text\":\"bash is forbidden\"}\\n'\n", encoding="utf-8")
            fake.chmod(fake.stat().st_mode | stat.S_IXUSR)
            result = OpenCodeEvalRunner(model="ollama/qwen3:14b", opencode=str(fake), timeout_seconds=5).run(workspace=root / "workspace")
            self.assertEqual(result.status, "FAILED")
            self.assertFalse(result.kslide_complete)
            self.assertIn("complete K-Slide run", result.reason or "")

    def test_numeric_value_in_wrong_region_fails(self) -> None:
        evidence = _evidence(
            regions=(EvidenceRegion("r1", selected_literal_candidate="매출 3.2조원", numeric_fact_ids=("r1-n01",)), EvidenceRegion("r2", selected_literal_candidate="영업이익")),
            numeric_facts=({"fact_id": "r1-n01", "source_object_id": "r1", "source_region_id": "r1", "source_string": "3.2조원", "raw_value": 3.2, "canonical_value": 3.2e12, "scale_factor": 1e12, "source_unit": "조원", "semantic_quantity": "currency", "currency": "KRW", "direction": None},),
            required_source_ids=("r1", "r2"),
        )
        score = score_translation_patch(type("Scenario", (), {"gold": {}})(), evidence, {"regions": [{"region_id": "r1", "english": "Revenue"}, {"region_id": "r2", "english": "Operating profit: KRW 3.2 trillion"}]})
        self.assertEqual(score["numeric_fidelity"], 0.0)

    def test_table_number_in_wrong_row_fails(self) -> None:
        table = EvidenceTable("t1", row_count=2, column_count=2, cells=(EvidenceTableCell("t1-r0-c0", 0, 0, source_text="매출", numeric_fact_ids=()), EvidenceTableCell("t1-r0-c1", 0, 1, source_text="3.2조원", numeric_fact_ids=("t1-r0-c1-n01",)), EvidenceTableCell("t1-r1-c0", 1, 0, source_text="영업이익", numeric_fact_ids=()), EvidenceTableCell("t1-r1-c1", 1, 1, source_text="500억원", numeric_fact_ids=("t1-r1-c1-n01",))))
        evidence = _evidence(tables=(table,), numeric_facts=({"fact_id": "t1-r0-c1-n01", "source_object_id": "t1-r0-c1", "source_cell_id": "t1-r0-c1", "source_string": "3.2조원", "raw_value": 3.2, "canonical_value": 3.2e12, "scale_factor": 1e12, "source_unit": "조원", "semantic_quantity": "currency", "currency": "KRW", "direction": None}, {"fact_id": "t1-r1-c1-n01", "source_object_id": "t1-r1-c1", "source_cell_id": "t1-r1-c1", "source_string": "500억원", "raw_value": 500, "canonical_value": 5e10, "scale_factor": 1e8, "source_unit": "억원", "semantic_quantity": "currency", "currency": "KRW", "direction": None}), required_source_ids=("t1",))
        patch = {"regions": [], "tables": [{"table_id": "t1", "cells": [{"cell_id": "t1-r0-c0", "english": "Revenue"}, {"cell_id": "t1-r0-c1", "english": "KRW 50 billion"}, {"cell_id": "t1-r1-c0", "english": "Operating profit"}, {"cell_id": "t1-r1-c1", "english": "KRW 3.2 trillion"}]}]}
        score = score_translation_patch(type("Scenario", (), {"gold": {}})(), evidence, patch)
        self.assertEqual(score["numeric_fidelity"], 0.0)

    def test_modality_correct_elsewhere_but_wrong_source_region_fails(self) -> None:
        evidence = _evidence(regions=(EvidenceRegion("r1", selected_literal_candidate="2H 적용 검토"), EvidenceRegion("r2", selected_literal_candidate="2H 적용 예정")), required_source_ids=("r1", "r2"))
        scenario = type("Scenario", (), {"gold": {"modality": {"source_text": "2H 적용 검토", "commitment": "under_review", "speech_act": "plan"}}})()
        patch = {"regions": [{"region_id": "r1", "english": "Deployment is scheduled", "commitment_status": "scheduled", "speech_act": "plan"}, {"region_id": "r2", "english": "Deployment is under review", "commitment_status": "under_review", "speech_act": "plan"}]}
        self.assertEqual(score_translation_patch(scenario, evidence, patch)["modality"], 0.0)

    def test_hangul_is_checked_in_table_and_claim_surfaces(self) -> None:
        patch = {"regions": [], "tables": [{"table_id": "t1", "cells": [{"cell_id": "c1", "english": "검토"}]}], "visual_interpretations": [], "executive_claims": [{"claim_id": "c", "kind": "takeaway", "text": "검토 필요", "evidence_ids": ["r1"]}]}
        score = score_translation_patch(type("Scenario", (), {"gold": {}})(), _evidence(), patch)
        self.assertIn("c1", score["hangul_by_surface"]["table_cells"])
        self.assertIn("c", score["hangul_by_surface"]["executive_claims"])

    def test_executive_claim_with_valid_evidence_but_wrong_semantics_fails(self) -> None:
        evidence = _evidence(regions=(EvidenceRegion("r1", selected_literal_candidate="2H 적용 검토"),), required_source_ids=("r1",))
        scenario = type("Scenario", (), {"gold": {"executive_claims": [{"kind": "decision_status", "acceptable_phrases": ["under review", "review"]}]}})()
        patch = {"regions": [{"region_id": "r1", "english": "Deployment remains under review", "commitment_status": "under_review"}], "executive_claims": [{"claim_id": "c1", "kind": "decision_status", "text": "Deployment is confirmed", "evidence_ids": ["r1"], "uncertainty": "low"}]}
        self.assertIn("WRONG_DECISION_STATUS_SEMANTICS", score_translation_patch(scenario, evidence, patch)["executive_claim_failures"])

    def test_champion_rejects_protected_regression_and_new_critical_type(self) -> None:
        baseline = {"critical_failure_count": 0, "critical_failure_types": [], "mean_scores": {"numeric_fidelity": 1.0}, "by_category": {"financial_table": {"numeric_fidelity": 1.0, "table_cell_fidelity": 1.0, "coverage": 1.0}}}
        candidate = {"critical_failure_count": 0, "critical_failure_types": ["NEW_CRITICAL"], "mean_scores": {"numeric_fidelity": 1.1}, "by_category": {"financial_table": {"numeric_fidelity": 0.9, "table_cell_fidelity": 1.0, "coverage": 1.0}}}
        comparison = compare_aggregate(baseline, candidate)
        self.assertFalse(comparison["accepted"])
        self.assertIn("financial_table", comparison["protected_regressions"])
        self.assertEqual(comparison["new_critical_types"], ["NEW_CRITICAL"])

    def test_multislide_deck_scenarios_cover_resume_sizes(self) -> None:
        decks = deck_scenarios()
        self.assertEqual([deck.slide_count for deck in decks], [3, 5, 10, 20, 50])
        self.assertTrue(all(len(deck.slides) == deck.slide_count for deck in decks))
        self.assertTrue(all(deck.gold["resume_safe"] for deck in decks))

    def test_deck_generator_emits_real_multislide_pptx_when_available(self) -> None:
        try:
            from evals.generator import generate_deck_pptx
            from pptx import Presentation
        except ImportError:
            self.skipTest("python-pptx is not installed")
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "linked-deck.pptx"
            deck = deck_scenarios()[0]
            self.assertTrue(generate_deck_pptx(deck.slides, destination))
            self.assertEqual(len(Presentation(str(destination)).slides), deck.slide_count)


if __name__ == "__main__":
    unittest.main()
