from __future__ import annotations

import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from k_slide.evidence_ir import load_evidence
from k_slide.extraction import extract_run
from k_slide.ingest import prepare_run
from k_slide.normalization import normalize_run
from k_slide.normalization import _pptx_native
from k_slide.queue import WorkUnitStatus, load_queue
from k_slide.state import RunPhase, load_state


@unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow is optional in the base development environment")
class ImageNormalizationTests(unittest.TestCase):
    def test_image_normalization_crops_and_builds_engine_evidence(self) -> None:
        from PIL import Image, ImageDraw

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image_path = root / "review.png"
            image = Image.new("RGB", (1200, 800), "white")
            ImageDraw.Draw(image).rectangle((100, 100, 1100, 700), outline="black", width=4)
            image.save(image_path)
            run = prepare_run(root, explicit_paths=[str(image_path)])
            result = normalize_run(run)
            self.assertEqual(len(result.documents), 1)
            self.assertEqual(load_state(run).phase, RunPhase.NORMALIZED)
            evidence_values = extract_run(run)
            self.assertEqual(load_state(run).phase, RunPhase.EXTRACTED)
            self.assertEqual(len(evidence_values), 1)
            evidence = load_evidence(run, "slide-001")
            self.assertEqual(evidence.evidence_revision, evidence.computed_revision())
            self.assertTrue(evidence.regions[0].crop_original_path)
            self.assertTrue((run / evidence.regions[0].crop_original_path).is_file())
            self.assertEqual(load_queue(run).work_units[0].status, WorkUnitStatus.READY)


@unittest.skipUnless(importlib.util.find_spec("fitz"), "PyMuPDF is optional in the base development environment")
class PDFNormalizationTests(unittest.TestCase):
    def test_pdf_pages_become_stable_units(self) -> None:
        import fitz

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf_path = root / "review.pdf"
            document = fitz.open()
            for index in range(3):
                page = document.new_page(width=400, height=300)
                page.insert_text((40, 60), f"Synthetic page {index + 1}")
            document.save(pdf_path)
            document.close()
            run = prepare_run(root, explicit_paths=[str(pdf_path)])
            result = normalize_run(run)
            self.assertEqual(len(result.documents[0].units), 3)
            self.assertEqual([unit.work_unit_id for unit in result.documents[0].units], ["page-001-0001", "page-001-0002", "page-001-0003"])


@unittest.skipUnless(importlib.util.find_spec("PIL") and importlib.util.find_spec("pptx"), "Pillow and python-pptx are optional in the base development environment")
class PPTXNativeExtractionTests(unittest.TestCase):
    def test_native_shapes_and_table_geometry_are_engine_evidence(self) -> None:
        from PIL import Image
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "review.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
            textbox.text = "검토 필요"
            table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(2))
            table_shape.table.cell(0, 0).text = "항목"
            table_shape.table.cell(0, 1).text = "수치"
            presentation.save(source)
            render = root / "slide-001.png"
            Image.new("RGB", (1920, 1080), "white").save(render)
            native = _pptx_native(source, root, "source-001", "doc-001", [render])

            self.assertEqual(len(native), 1)
            self.assertEqual((native[0].width_px, native[0].height_px), (1920, 1080))
            objects = json.loads((root / native[0].native_evidence_path).read_text())["objects"]
            self.assertEqual(len(objects), 2)
            self.assertEqual(objects[0]["bbox_px"], [192, 144, 768, 288])
            self.assertEqual(objects[1]["table"]["row_count"], 2)
            self.assertEqual(objects[1]["table"]["column_count"], 2)


if __name__ == "__main__":
    unittest.main()
