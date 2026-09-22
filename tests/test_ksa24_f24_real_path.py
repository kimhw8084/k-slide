from __future__ import annotations

import copy
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from k_slide.cli import _conflict_assess, _next, _submit
from k_slide.evidence_ir import load_evidence
from k_slide.errors import KSlideError
from k_slide.ingest import prepare_run
from k_slide.io import atomic_write_json, read_json
from k_slide.rendering.reports import render_run
from k_slide.storage import StorageArtifact, storage_path
from k_slide.terminology import Termbase
from k_slide.verify import VerificationResult, _validate_slide, finalize_run, verify_run
from tests.reference_fixtures import reference_environment


def _add_arrowhead(connector: object, element_name: str, arrow_type: str) -> None:
    from pptx.oxml.xmlchemy import OxmlElement

    line = connector._element.spPr.get_or_add_ln()  # type: ignore[attr-defined]
    arrow = OxmlElement(f"a:{element_name}")
    arrow.set("type", arrow_type)
    line.append(arrow)


def _representative_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    mixed = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(3), Inches(0.4))
    mixed.text = "AI Platform 검토 중"
    decision = slide.shapes.add_textbox(Inches(0.5), Inches(0.8), Inches(3), Inches(0.4))
    decision.text = "출시 결정"

    group = slide.shapes.add_group_shape()
    group.left = Inches(0.5)
    group.top = Inches(2)
    group.width = Inches(4)
    group.height = Inches(1)
    start = group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0), Inches(0), Inches(1), Inches(0.5))
    start.text = "A"
    end = group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2), Inches(0), Inches(1), Inches(0.5))
    end.text = "B"

    forward = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.5), Inches(2.25), Inches(3), Inches(2.25))
    forward.begin_connect(start, 0)
    forward.end_connect(end, 0)
    _add_arrowhead(forward, "tailEnd", "triangle")

    reverse_arrow = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.5), Inches(2.5), Inches(3), Inches(2.5))
    reverse_arrow.begin_connect(start, 0)
    reverse_arrow.end_connect(end, 0)
    _add_arrowhead(reverse_arrow, "headEnd", "triangle")

    undirected = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.5), Inches(2.75), Inches(3), Inches(2.75))
    undirected.begin_connect(start, 0)
    undirected.end_connect(end, 0)

    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2", "Q3", "Q4"]
    chart_data.add_series("North", (1.0, 2.0, 3.0, 4.0))
    chart_data.add_series("South", (4.0, 3.0, 2.0, 1.0))
    chart_data.add_series("Nullable", (None, 2.5, None, -1.0))
    chart_data.add_series("Formatted", (1234.5, -1234.5, 0.0, 2.25))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(5), Inches(0.5), Inches(4), Inches(3), chart_data).chart
    chart.value_axis.tick_labels.number_format = "0.0%;-0.0%;-"

    table_shape = slide.shapes.add_table(3, 3, Inches(0.5), Inches(4), Inches(5), Inches(2))
    table = table_shape.table
    table._tbl.tblPr.set("firstRow", "1")
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "USD millions"
    table.cell(0, 2).text = "Metric"
    table.cell(1, 0).text = "1.5"
    table.cell(1, 1).text = "Revenue"
    table.cell(1, 2).text = "   "
    table.cell(2, 0).text = "2.25"
    table.cell(2, 1).text = "AI Platform 검토 중"
    table.cell(2, 2).text = "7"

    presentation.save(path)


def _fixture_render(_source: Path, run_dir: Path, document_id: str) -> list[Path]:
    # The fixture is parsed by python-pptx and the normal extractor; this only
    # replaces the optional office renderer, which is unavailable in CI.
    from PIL import Image

    render = storage_path(run_dir, StorageArtifact.NORMALIZED_RENDER, f"normalized/{document_id}-slide-0001.png", create_parent=True)
    Image.new("RGB", (1600, 900), "white").save(render)
    render.chmod(0o600)
    return [render]


def _region_patch(evidence: object, region: object) -> dict[str, object]:
    source = region.selected_literal_candidate
    value: dict[str, object] = {
        "region_id": region.region_id,
        "english": "" if source is None else source,
        "term_ids": [],
        "unresolved": False,
    }
    if source == "AI Platform 검토 중":
        value.update({"english": "AI Platform under review", "commitment_status": "under_review", "speech_act": "status", "evidence_ids": [region.region_id]})
    elif source == "출시 결정":
        value.update({"english": "Launch decided", "commitment_status": "decided", "speech_act": "decision", "evidence_ids": [region.region_id]})
    return value


def _valid_patch(evidence: object) -> dict[str, object]:
    chart = next(item for item in evidence.visual_elements if item.get("kind") == "chart")
    chart_id = str(chart["element_id"])
    chart_data = chart["chart"]
    categories = chart_data["categories"]
    connectors = [item for item in evidence.visual_elements if item.get("kind") == "connector"]
    forward = next(item for item in connectors if item["connector"]["direction_evidence"] == "start_to_end")
    connector = forward["connector"]
    forward_relation = {
        "relation_id": "real-forward",
        "interpretation": "A flows to B",
        "evidence_ids": [forward["element_id"]],
        "source_element_ids": [connector["from_element_id"], connector["to_element_id"], forward["element_id"]],
        "relation_type": "next",
        "direction": "left_to_right",
        "provenance": "source_fact",
    }
    chart_relations = [
        {
            "relation_id": "real-north-trend",
            "interpretation": "North increases",
            "evidence_ids": [chart_id],
            "source_element_ids": [chart_id],
            "provenance": "source_fact",
            "chart_claim": {"chart_element_id": chart_id, "kind": "trend", "series_index": 0, "series_name": "North", "direction": "increasing"},
        },
        {
            "relation_id": "real-south-trend",
            "interpretation": "South decreases",
            "evidence_ids": [chart_id],
            "source_element_ids": [chart_id],
            "provenance": "source_fact",
            "chart_claim": {"chart_element_id": chart_id, "kind": "trend", "series_index": 1, "series_name": "South", "direction": "decreasing"},
        },
        {
            "relation_id": "real-null-point",
            "interpretation": "Nullable Q1 is blank",
            "evidence_ids": [chart_id],
            "source_element_ids": [chart_id],
            "provenance": "source_fact",
            "chart_claim": {"chart_element_id": chart_id, "kind": "point_value", "series_index": 2, "series_name": "Nullable", "point_index": 0, "category": categories[0], "is_blank": True},
        },
        {
            "relation_id": "real-negative-value",
            "interpretation": "Formatted is -1234.5",
            "evidence_ids": [chart_id],
            "source_element_ids": [chart_id],
            "provenance": "source_fact",
            "chart_claim": {"chart_element_id": chart_id, "kind": "point_value", "series_index": 3, "series_name": "Formatted", "point_index": 1, "category": categories[1], "value": -1234.5, "is_blank": False},
        },
        {
            "relation_id": "real-ranking",
            "interpretation": "Formatted is lowest at Q2",
            "evidence_ids": [chart_id],
            "source_element_ids": [chart_id],
            "provenance": "source_fact",
            "chart_claim": {"chart_element_id": chart_id, "kind": "ranking", "series_index": 3, "series_name": "Formatted", "point_index": 1, "category": categories[1], "ranking": "lowest", "rank": 1},
        },
        {
            "relation_id": "real-comparison",
            "interpretation": "North is below South at Q1",
            "evidence_ids": [chart_id],
            "source_element_ids": [chart_id],
            "provenance": "source_fact",
            "chart_claim": {"chart_element_id": chart_id, "kind": "comparison", "series_index": 0, "series_name": "North", "point_index": 0, "category": categories[0], "other_series_index": 1, "other_series_name": "South", "operator": "less_than"},
        },
    ]
    table = evidence.tables[0]
    cells = []
    for cell in table.cells:
        source = cell.source_text
        english = "" if cell.cell_state in {"blank", "merge_continuation"} else ("AI Platform under review 검토 중" if source == "AI Platform 검토 중" else source or "")
        cell_patch: dict[str, object] = {"cell_id": cell.cell_id, "english": english, "unresolved": False}
        if source == "AI Platform 검토 중":
            cell_patch.update({
                "commitment_status": "under_review",
                "speech_act": "status",
                "evidence_ids": [cell.cell_id],
                "hangul_retention": {"reason": "Retain the source label alongside its translation.", "evidence_id": cell.cell_id},
            })
        cells.append(cell_patch)
    under_review_region = next(region for region in evidence.regions if region.selected_literal_candidate == "AI Platform 검토 중")
    claim = {
        "claim_id": "real-status-claim",
        "kind": "decision_status",
        "text": "AI Platform remains under review; source wording: 검토 중",
        "evidence_ids": [under_review_region.region_id],
        "uncertainty": "low",
        "provenance": "supported_interpretation",
        "hangul_retention": {"reason": "Retain the source wording alongside the translated status.", "evidence_id": under_review_region.region_id},
    }
    return {
        "schema_version": "1.0",
        "work_unit_id": evidence.work_unit_id,
        "evidence_revision": evidence.evidence_revision,
        "regions": [_region_patch(evidence, region) for region in evidence.regions],
        "tables": [{"table_id": table.table_id, "cells": cells}],
        "visual_interpretations": [forward_relation, *chart_relations],
        "executive_claims": [claim],
    }


@unittest.skipUnless(importlib.util.find_spec("PIL") and importlib.util.find_spec("pptx"), "Pillow and python-pptx are required for the real-path fixture")
class KSA24F24RealPathTests(unittest.TestCase):
    def test_generated_pptx_survives_employee_runtime_and_rejects_adversaries(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "ksa24-representative.pptx"
            _representative_pptx(source)
            environment = reference_environment("ksa24-f24-real-path")
            with patch("k_slide.normalization._render_pptx", side_effect=_fixture_render):
                run = prepare_run(root, explicit_paths=[str(source)], perform_processing=True, environment_identity=environment)

            next_value = _next(root, run.name, None, environment)
            self.assertEqual(next_value["status"], "READY")
            evidence = load_evidence(run, next_value["work_unit_id"])
            self.assertEqual(evidence.source["source_language_policy"], "unicode-script-v1")
            self.assertIn("AI Platform", {span for region in evidence.regions for span in region.source_english_spans})
            self.assertEqual({region.language for region in evidence.regions if region.selected_literal_candidate == "AI Platform 검토 중"}, {"mixed"})

            table = evidence.tables[0]
            self.assertIsNotNone(table)
            self.assertEqual(table.unit, "USD millions")
            self.assertIn("merge_continuation", {cell.cell_state for cell in table.cells})
            self.assertIn("blank", {cell.cell_state for cell in table.cells})
            self.assertIn("nonblank", {cell.cell_state for cell in table.cells})
            protected_cell = next(cell for cell in table.cells if cell.source_text == "AI Platform 검토 중")
            self.assertEqual(protected_cell.source_language, "mixed")
            self.assertIn("AI Platform", protected_cell.source_english_spans)
            numeric = {fact.get("source_cell_id"): fact for fact in evidence.numeric_facts}
            self.assertEqual(numeric[next(cell.cell_id for cell in table.cells if cell.source_text == "1.5")]["canonical_value"], 1.5 * 10**6)

            chart = next(item for item in evidence.visual_elements if item.get("kind") == "chart")
            series = chart["chart"]["series"]
            self.assertEqual([point["value"] for point in series[0]["points"]], [1.0, 2.0, 3.0, 4.0])
            self.assertEqual([point["value"] for point in series[1]["points"]], [4.0, 3.0, 2.0, 1.0])
            self.assertTrue(series[2]["points"][0]["is_blank"])
            self.assertEqual(series[3]["points"][1]["value"], -1234.5)
            self.assertEqual(chart["chart"]["unit_labels"]["value"], "0.0%;-0.0%;-")

            connectors = [item for item in evidence.visual_elements if item.get("kind") == "connector"]
            self.assertEqual({item["connector"]["direction_evidence"] for item in connectors}, {"start_to_end", "end_to_start", "undirected"})
            self.assertTrue(all("-shape-003-" in item["connector"]["from_element_id"] for item in connectors))

            valid = _valid_patch(evidence)
            chart_id = chart["element_id"]
            forward = next(item for item in connectors if item["connector"]["direction_evidence"] == "start_to_end")
            connector = forward["connector"]

            adversaries = []
            fabricated = copy.deepcopy(valid)
            fabricated["visual_interpretations"] = [{
                "relation_id": "fabricated",
                "interpretation": "invented",
                "evidence_ids": ["fabricated-element"],
                "source_element_ids": ["fabricated-element"],
                "provenance": "supported_interpretation",
            }]
            adversaries.append((fabricated, "KSLIDE_UNKNOWN_SOURCE_ELEMENT"))

            reversed_connector = copy.deepcopy(valid)
            reversed_connector["visual_interpretations"] = [{
                "relation_id": "reversed",
                "interpretation": "B flows to A",
                "evidence_ids": [forward["element_id"]],
                "source_element_ids": [connector["to_element_id"], connector["from_element_id"], forward["element_id"]],
                "relation_type": "next",
                "direction": "right_to_left",
                "provenance": "source_fact",
            }]
            adversaries.append((reversed_connector, "KSLIDE_SCHEMA_INVALID"))

            undirected_connector = copy.deepcopy(valid)
            undirected = next(item for item in connectors if item["connector"]["direction_evidence"] == "undirected")
            undirected_connector["visual_interpretations"] = [{
                "relation_id": "undirected-source-fact",
                "interpretation": "A flows to B",
                "evidence_ids": [undirected["element_id"]],
                "source_element_ids": [undirected["connector"]["from_element_id"], undirected["connector"]["to_element_id"], undirected["element_id"]],
                "relation_type": "next",
                "direction": "left_to_right",
                "provenance": "source_fact",
            }]
            adversaries.append((undirected_connector, "KSLIDE_SCHEMA_INVALID"))

            wrong_trend = copy.deepcopy(valid)
            wrong_trend["visual_interpretations"] = [{
                "relation_id": "reversed-trend",
                "interpretation": "North decreases",
                "evidence_ids": [chart_id],
                "source_element_ids": [chart_id],
                "provenance": "source_fact",
                "chart_claim": {"chart_element_id": chart_id, "kind": "trend", "series_index": 0, "series_name": "North", "direction": "decreasing"},
            }]
            adversaries.append((wrong_trend, "KSLIDE_CHART_INTERPRETATION_MISMATCH"))

            blank_as_number = copy.deepcopy(valid)
            blank_as_number["visual_interpretations"] = [{
                "relation_id": "blank-as-number",
                "interpretation": "Nullable Q1 is zero",
                "evidence_ids": [chart_id],
                "source_element_ids": [chart_id],
                "provenance": "source_fact",
                "chart_claim": {"chart_element_id": chart_id, "kind": "point_value", "series_index": 2, "series_name": "Nullable", "point_index": 0, "category": "Q1", "value": 0, "is_blank": False},
            }]
            adversaries.append((blank_as_number, "KSLIDE_CHART_INTERPRETATION_MISMATCH"))

            strengthened = copy.deepcopy(valid)
            under_review = next(region for region in evidence.regions if region.selected_literal_candidate == "AI Platform 검토 중")
            strengthened["regions"] = [
                {**region, "english": "AI Platform approved", "commitment_status": "under_review", "speech_act": "status", "evidence_ids": [under_review.region_id]}
                if region["region_id"] == under_review.region_id else region
                for region in strengthened["regions"]
            ]
            adversaries.append((strengthened, "KSLIDE_MODALITY_MISMATCH"))

            table_cell_patch = next(cell for cell in valid["tables"][0]["cells"] if cell["cell_id"] == protected_cell.cell_id)
            omitted_cell_status = copy.deepcopy(valid)
            omitted_cell = next(cell for cell in omitted_cell_status["tables"][0]["cells"] if cell["cell_id"] == protected_cell.cell_id)
            omitted_cell.pop("commitment_status")
            omitted_cell.pop("speech_act")
            adversaries.append((omitted_cell_status, "KSLIDE_MODALITY_MISMATCH"))

            mismatched_cell_status = copy.deepcopy(valid)
            next(cell for cell in mismatched_cell_status["tables"][0]["cells"] if cell["cell_id"] == protected_cell.cell_id).update({"commitment_status": "decided", "speech_act": "decision"})
            adversaries.append((mismatched_cell_status, "KSLIDE_MODALITY_MISMATCH"))

            strengthened_cell = copy.deepcopy(valid)
            next(cell for cell in strengthened_cell["tables"][0]["cells"] if cell["cell_id"] == protected_cell.cell_id)["english"] = "AI Platform approved"
            adversaries.append((strengthened_cell, "KSLIDE_MODALITY_MISMATCH"))

            weakened_cell = copy.deepcopy(valid)
            next(cell for cell in weakened_cell["tables"][0]["cells"] if cell["cell_id"] == protected_cell.cell_id)["english"] = "AI Platform status"
            adversaries.append((weakened_cell, "KSLIDE_MODALITY_MISMATCH"))

            approved_claim = copy.deepcopy(valid)
            approved_claim["executive_claims"][0]["text"] = "AI Platform approved"
            adversaries.append((approved_claim, "KSLIDE_MODALITY_MISMATCH"))

            hangul_cell = copy.deepcopy(valid)
            next(cell for cell in hangul_cell["tables"][0]["cells"] if cell["cell_id"] == protected_cell.cell_id)["english"] = "AI Platform under review 검토 중"
            next(cell for cell in hangul_cell["tables"][0]["cells"] if cell["cell_id"] == protected_cell.cell_id).pop("hangul_retention")
            adversaries.append((hangul_cell, "KSLIDE_REQUIRED_ENGLISH"))

            hangul_visual = copy.deepcopy(valid)
            hangul_visual["visual_interpretations"][0]["interpretation"] = "A flows to B 한글"
            adversaries.append((hangul_visual, "KSLIDE_REQUIRED_ENGLISH"))

            hangul_claim = copy.deepcopy(valid)
            hangul_claim["executive_claims"][0]["text"] = "AI Platform remains under review 검토 중"
            hangul_claim["executive_claims"][0].pop("hangul_retention")
            adversaries.append((hangul_claim, "KSLIDE_REQUIRED_ENGLISH"))

            for payload, expected_code in adversaries:
                with self.subTest(expected_code=expected_code):
                    with self.assertRaises(KSlideError) as raised:
                        _submit(root, run.name, json.dumps(payload, ensure_ascii=False), None, environment)
                    self.assertEqual(getattr(raised.exception, "code", None).value, expected_code)

            accepted = _submit(root, run.name, json.dumps(valid, ensure_ascii=False), None, environment)
            self.assertEqual(accepted["status"], "ACCEPTED")
            self.assertEqual(table_cell_patch["english"], "AI Platform under review 검토 중")
            self.assertTrue(storage_path(run, StorageArtifact.REPORT, "05_final_report.md").is_file())
            report = storage_path(run, StorageArtifact.REPORT, "05_final_report.md").read_text(encoding="utf-8")
            self.assertIn("USD millions", report)
            self.assertIn("AI Platform 검토 중", report)  # Raw source display remains intact.

            self.assertEqual(_next(root, run.name, None, environment)["status"], "CONFLICT_ASSESSMENT_REQUIRED")
            assessed = _conflict_assess(root, run.name, json.dumps({"schema_version": "1.0", "candidate_groups": []}), None, environment)
            self.assertEqual(assessed["status"], "ASSESSED_ZERO_CONFLICTS")
            self.assertEqual(_next(root, run.name, None, environment)["status"], "ALL_TRANSLATED")

            render_run(run)
            canonical_path = storage_path(run, StorageArtifact.CANONICAL_IR, f"ir/{next_value['work_unit_id']}.json")
            canonical = read_json(canonical_path)
            canonical_cell = next(cell for cell in canonical["tables"][0]["cells"] if cell["cell_id"] == protected_cell.cell_id)
            self.assertEqual(canonical_cell["hangul_retention"]["evidence_id"], protected_cell.cell_id)
            self.assertEqual(canonical["executive_semantics"]["executive_claims"][0]["hangul_retention"]["evidence_id"], under_review.region_id)

            def tamper_unretained_hangul(value: dict[str, object]) -> None:
                cell = next(cell for cell in value["tables"][0]["cells"] if cell["cell_id"] == protected_cell.cell_id)  # type: ignore[index]
                cell.update({"translation": "AI Platform under review 검토 중"})
                cell.pop("hangul_retention", None)

            for mutate, expected_code in (
                (
                    lambda value: next(cell for cell in value["tables"][0]["cells"] if cell["cell_id"] == protected_cell.cell_id).update({"commitment_status": "invented"}),
                    "KSLIDE_SCHEMA_INVALID",
                ),
                (
                    lambda value: next(cell for cell in value["tables"][0]["cells"] if cell["cell_id"] == protected_cell.cell_id).update({"commitment_status": "decided", "speech_act": "decision"}),
                    "KSLIDE_MODALITY_MISMATCH",
                ),
                (
                    lambda value: next(cell for cell in value["tables"][0]["cells"] if cell["cell_id"] == protected_cell.cell_id).update({"translation": "AI Platform approved"}),
                    "KSLIDE_MODALITY_MISMATCH",
                ),
                (
                    tamper_unretained_hangul,
                    "KSLIDE_REQUIRED_ENGLISH",
                ),
                (
                    lambda value: value["executive_semantics"]["executive_claims"][0].update({"text": "AI Platform approved"}),
                    "KSLIDE_MODALITY_MISMATCH",
                ),
            ):
                tampered = copy.deepcopy(canonical)
                mutate(tampered)
                atomic_write_json(canonical_path, tampered)
                result = VerificationResult(status="PASS", run_id=run.name)
                _validate_slide(run, next_value["work_unit_id"], result, termbase=Termbase("1.0", ()))
                self.assertIn(expected_code, {issue.code for issue in result.issues})
            atomic_write_json(canonical_path, canonical)

            verified = verify_run(run, environment_identity=environment)
            self.assertTrue(verified.passed, verified.as_dict())
            finalized = finalize_run(run, environment_identity=environment)
            self.assertTrue(finalized.passed, finalized.as_dict())
            self.assertTrue((run / "RUN_COMPLETE.md").is_file())


if __name__ == "__main__":
    unittest.main()
