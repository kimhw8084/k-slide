"""Generate real, deterministic, non-confidential evaluation artifacts."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable

from .scenarios import Scenario


def _font(size: int):
    from PIL import ImageFont

    candidates = (
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Noto Sans CJK KR.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
    )
    for candidate in candidates:
        if Path(candidate).is_file():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def generate_image(scenario: Scenario, destination: Path, *, width: int = 1600, height: int = 900, image_format: str = "PNG") -> Path:
    from PIL import Image, ImageDraw

    destination.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (width, height), (247, 249, 252))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 110), fill=(16, 49, 92))
    draw.text((60, 30), scenario.title_ko, font=_font(46), fill="white")
    y = 170
    for index, line in enumerate(scenario.body_ko, start=1):
        draw.rounded_rectangle((70, y - 12, width - 70, y + 74), radius=12, fill=(255, 255, 255), outline=(190, 202, 219), width=2)
        draw.text((100, y + 10), f"{index}. {line}", font=_font(30), fill=(26, 33, 45))
        y += 115
    draw.text((70, height - 60), f"Synthetic evaluation case: {scenario.scenario_id} | {scenario.category}", font=_font(20), fill=(90, 100, 115))
    image.save(destination, format=image_format, optimize=True)
    return destination


def _generate_pdf(source_png: Path, destination: Path) -> bool:
    try:
        import fitz
        from PIL import Image
    except ImportError:
        return False
    with fitz.open() as document:
        with Image.open(source_png) as image:
            page = document.new_page(width=image.width, height=image.height)
        page.insert_image(page.rect, filename=str(source_png))
        document.save(destination)
    return True


def _generate_pptx(scenario: Scenario, destination: Path) -> bool:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        return False
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.7)).text_frame
    title.text = scenario.title_ko
    title.paragraphs[0].font.size = Pt(28)
    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(11.8), Inches(4.8)).text_frame
    body.clear()
    for index, line in enumerate(scenario.body_ko):
        paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(20)
        paragraph.level = 0
    presentation.save(destination)
    return True


def generate_artifacts(scenarios: Iterable[Scenario], output: Path, *, formats: tuple[str, ...] = ("png",)) -> dict[str, int]:
    output.mkdir(parents=True, exist_ok=True)
    counts = {name: 0 for name in formats}
    for scenario in scenarios:
        case_dir = output / scenario.scenario_id
        case_dir.mkdir(parents=True, exist_ok=True)
        (case_dir / "gold.json").write_text(json.dumps(scenario.as_dict(), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        for format_name in formats:
            suffix = format_name.lower()
            destination = case_dir / f"source.{suffix}"
            if suffix in {"png", "jpg", "jpeg", "webp"}:
                generate_image(scenario, destination, image_format="JPEG" if suffix in {"jpg", "jpeg"} else suffix.upper())
                counts[format_name] += 1
            elif suffix == "pdf":
                temporary_png = case_dir / "_source_for_pdf.png"
                generate_image(scenario, temporary_png)
                if _generate_pdf(temporary_png, destination):
                    counts[format_name] += 1
                temporary_png.unlink(missing_ok=True)
            elif suffix == "pptx":
                if _generate_pptx(scenario, destination):
                    counts[format_name] += 1
    return counts
