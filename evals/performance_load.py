"""Bounded K-Slide product-path performance and admission measurement."""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import math
import os
import platform
import shutil
import subprocess
import sys
import tempfile
import time
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from k_slide.errors import KSlideError
from k_slide.environment import RunEnvironmentIdentity
from k_slide import EXECUTION_CONTRACT_VERSION
from k_slide.paas import (
    AuthorizedScopeContext,
    PaaSController,
    PaaSJobRequest,
    ReferencePaaSJobService,
    RuntimeIdentity,
    ScopedAdmissionPolicy,
)
from k_slide.resource_budget import ResourceBudget


BASE_SHA = "71d3b0a56eee84716a37712440dd1763fd078c66"
REQUEST_ID = "chg16-performance-load-budget-01"
MAX_REPETITIONS = 5
MAX_WARMUPS = 1
_KNOWN_STDOUT_LIBRARY_WARNING = "warning: The `fitz` API is deprecated and will be removed in future. Use `import pymupdf` instead."
CORE_TOPOLOGIES = (
    ("local_cli_reference", "explicit_paths"),
    ("opencode_shared_core_reference", "attachment_reference"),
    ("cloud_vscode_shared_cli_reference", "workspace_file_reference"),
)


@dataclass(frozen=True)
class Scenario:
    scenario_id: str
    scenario_class: str
    format: str
    pages: int
    width_px: int
    height_px: int


SCENARIOS = (
    Scenario("screenshot-1", "screenshot", "png", 1, 1600, 900),
    Scenario("presentation-16-pdf", "presentation_10_20_slides", "pdf", 16, 1600, 900),
    Scenario("presentation-55-pdf", "presentation_50_plus_slides", "pdf", 55, 1600, 900),
    Scenario("presentation-16-pptx", "presentation_10_20_slides", "pptx", 16, 1600, 900),
    Scenario("presentation-55-pptx", "presentation_50_plus_slides", "pptx", 55, 1600, 900),
)


def _git(root: Path, *args: str) -> str:
    result = subprocess.run(["git", *args], cwd=root, check=True, capture_output=True, text=True)
    return result.stdout.strip()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _environment() -> dict[str, Any]:
    libraries: dict[str, str] = {}
    for package in ("Pillow", "PyMuPDF", "python-pptx", "paddleocr", "paddlepaddle"):
        try:
            libraries[package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            libraries[package] = "NOT_INSTALLED"
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    opencode = shutil.which("opencode")
    return {
        "python_version": platform.python_version(),
        "python_implementation": platform.python_implementation(),
        "platform": platform.platform(),
        "system": platform.system(),
        "machine": platform.machine(),
        "processor": platform.processor() or "UNREPORTED",
        "cpu_count": os.cpu_count(),
        "libraries": libraries,
        "soffice_path_present": bool(soffice),
        "opencode_path_present": bool(opencode),
        "target_model": "google/gemma-4-31b-it",
        "target_model_status": "NOT_MEASURED_EXTERNAL_PROVIDER_UNAVAILABLE",
        "company_paas_status": "NOT_MEASURED_EXTERNAL_JOB_SERVICE_UNAVAILABLE",
    }


def _environment_identity(environment: dict[str, Any]) -> str:
    encoded = json.dumps(environment, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _write_scenario(path: Path, scenario: Scenario) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if scenario.format == "png":
        from PIL import Image, ImageDraw

        image = Image.new("RGB", (scenario.width_px, scenario.height_px), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((0, 0, scenario.width_px, 104), fill="#153b67")
        draw.text((64, 36), "Synthetic operating review", fill="white")
        for row in range(9):
            y = 160 + row * 72
            draw.line((64, y, 1520, y), fill="#ccd3dc", width=2)
            draw.text((80, y + 18), f"Metric {row + 1}: plan, result, next action", fill="#202a35")
        image.save(path, format="PNG", optimize=True)
        return
    if scenario.format == "pptx":
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches, Pt

        presentation = Presentation()
        presentation.slide_width = Inches(13.333333333333334)
        presentation.slide_height = Inches(7.5)
        for slide_index in range(scenario.pages):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333333333333334), Inches(0.9))
            header.fill.solid()
            header.fill.fore_color.rgb = RGBColor(21, 59, 103)
            header.line.fill.background()
            title = slide.shapes.add_textbox(Inches(0.45), Inches(0.15), Inches(12.2), Inches(0.45)).text_frame
            title.text = f"Synthetic review — slide {slide_index + 1}"
            title.paragraphs[0].font.size = Pt(20)
            for row in range(8):
                text = slide.shapes.add_textbox(Inches(0.6), Inches(1.25 + row * 0.65), Inches(12.0), Inches(0.35)).text_frame
                text.text = f"Workstream {row + 1} | planned action | current status | owner"
                text.paragraphs[0].font.size = Pt(11)
        presentation.save(path)
        return
    import fitz

    document = fitz.open()
    for page_index in range(scenario.pages):
        page = document.new_page(width=960, height=540)
        page.draw_rect(fitz.Rect(0, 0, 960, 65), color=(0.08, 0.23, 0.4), fill=(0.08, 0.23, 0.4))
        page.insert_text((34, 42), f"Synthetic review — slide {page_index + 1}", fontsize=19, color=(1, 1, 1))
        for row in range(8):
            y = 110 + row * 48
            page.draw_line((34, y + 12), (928, y + 12), color=(0.78, 0.81, 0.85), width=0.7)
            page.insert_text((44, y), f"Workstream {row + 1} | planned action | current status | owner", fontsize=10, color=(0.12, 0.16, 0.2))
    document.set_metadata({"title": "Synthetic K-Slide load scenario", "author": "K-Slide reference harness"})
    document.save(path, garbage=4, deflate=True)
    document.close()


def _percentile(samples: list[float], percentile: float) -> float | None:
    if not samples:
        return None
    ordered = sorted(samples)
    return ordered[max(0, math.ceil(percentile * len(ordered)) - 1)]


def _parse_product_cli_json(stdout: str) -> dict[str, Any] | None:
    """Parse one product CLI response, allowing only the known PyMuPDF warning."""

    try:
        value = json.loads(stdout)
    except json.JSONDecodeError:
        start = stdout.find("{")
        if start < 0:
            return None
        prefix = stdout[:start].strip()
        if not prefix or any(line != _KNOWN_STDOUT_LIBRARY_WARNING for line in prefix.splitlines()):
            return None
        try:
            value, end = json.JSONDecoder().raw_decode(stdout, start)
        except json.JSONDecodeError:
            return None
        if stdout[end:].strip():
            return None
    return value if isinstance(value, dict) else None


def _reference_environment(source_revision: str) -> RunEnvironmentIdentity:
    return RunEnvironmentIdentity.legacy_reference(
        runtime_ref="ksa36-local-reference-runtime",
        model_identity="NOT_INVOKED",
        ocr_identity="native-extraction-only",
        termbase_identity="core-reference-termbase",
        engine_contract_version=EXECUTION_CONTRACT_VERSION,
        source_revision=source_revision,
    )


def _measure_core_once(project_root: Path, workspace_root: Path, source: Path, topology: str, input_mode: str, environment: RunEnvironmentIdentity) -> dict[str, Any]:
    host_adapter = {
        "local_cli_reference": "local_cli",
        "opencode_shared_core_reference": "opencode",
        "cloud_vscode_shared_cli_reference": "cloud_vscode",
    }[topology]
    command = ["prepare", "--root", str(workspace_root), "--json", "--host-adapter", host_adapter]
    if input_mode == "explicit_paths":
        command.append(str(source))
    else:
        source_kind = "attachment" if input_mode == "attachment_reference" else "workspace_file"
        invocation = {
            "schema_version": "1.0",
            "adapter_version": "1.0",
            "input_refs": [{
                "source_kind": source_kind,
                "logical_name": source.name,
                "locator": str(source),
                "classification": "company_confidential",
            }],
        }
        command.extend(["--host-inputs-json", json.dumps(invocation, separators=(",", ":")), "--host-worktree", str(workspace_root)])

    runner = """import sys
from unittest.mock import patch
from k_slide import cli
from k_slide.environment import RunEnvironmentIdentity
real_prepare = cli.prepare_run
reference_environment = RunEnvironmentIdentity.legacy_reference(
    runtime_ref="ksa36-local-reference-runtime",
    model_identity="NOT_INVOKED",
    ocr_identity="native-extraction-only",
    termbase_identity="core-reference-termbase",
    source_revision=sys.argv[1],
)
def prepare_with_reference(root, **kwargs):
    return real_prepare(root, environment_identity=reference_environment, **kwargs)
with patch.object(cli, "prepare_run", side_effect=prepare_with_reference):
    raise SystemExit(cli.main(sys.argv[2:]))
"""
    process_environment = os.environ.copy()
    process_environment["PYTHONPATH"] = os.pathsep.join((str(project_root / "src"), str(project_root)))
    started = time.perf_counter_ns()
    completed = subprocess.run(
        [sys.executable, "-c", runner, environment.kslide_source_revision, *command],
        cwd=project_root,
        env=process_environment,
        capture_output=True,
        text=True,
        check=False,
    )
    elapsed_ms = (time.perf_counter_ns() - started) / 1_000_000
    output = _parse_product_cli_json(completed.stdout)
    if output is None:
        return {
            "status": "NOT_MEASURED",
            "process_exit_code": completed.returncode,
            "output_parse": "FAILED",
            "stage_timings_ms": {"product_cli_prepare_end_to_end": elapsed_ms},
        }
    phase = output.get("phase") or output.get("status")
    measured = completed.returncode == 0 and phase not in {"FAILED", "FAILED_INPUT", "FAILED_NORMALIZATION", "FAILED_EXTRACTION", "FAILED_RUNTIME"}
    queue = output.get("work_queue") if isinstance(output.get("work_queue"), dict) else {}
    counts = queue.get("counts") if isinstance(queue.get("counts"), dict) else {}
    return {
        "status": "MEASURED_REFERENCE_PRODUCT_PATH" if measured else "NOT_MEASURED",
        "process_exit_code": completed.returncode,
        "phase": phase,
        "error_code": output.get("error_code"),
        "stage_timings_ms": {"product_cli_prepare_end_to_end": elapsed_ms},
        "work_units": queue.get("total", counts.get("total")),
        "input_count": output.get("input_count"),
        "excluded_stages": ["OpenCode plugin lifecycle and actual provider call", "Gemma invocation and provider usage", "TranslationPatch generation", "deterministic verification", "finalization"],
    }


def _core_results(case_root: Path, scenario: Scenario, source_template: Path, warmups: int, repetitions: int, environment: RunEnvironmentIdentity) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    for topology, input_mode in CORE_TOPOLOGIES:
        all_samples: list[dict[str, Any]] = []
        for index in range(warmups + repetitions):
            measured = index >= warmups
            with tempfile.TemporaryDirectory(prefix="k-slide-ksa36-run-") as temporary:
                run_root = Path(temporary)
                input_dir = run_root / ".bench-input"
                input_dir.mkdir()
                source = input_dir / source_template.name
                shutil.copyfile(source_template, source)
                sample = _measure_core_once(Path(__file__).resolve().parents[1], run_root, source, topology, input_mode, environment)
            if measured:
                sample["repetition"] = index - warmups + 1
                all_samples.append(sample)
        durations = [sample["stage_timings_ms"]["product_cli_prepare_end_to_end"] for sample in all_samples if sample.get("status") == "MEASURED_REFERENCE_PRODUCT_PATH"]
        if scenario.format in {"pdf", "pptx"}:
            render_width = math.ceil(960 * 220 / 72)
            render_height = math.ceil(540 * 220 / 72)
            expected_pixels = render_width * render_height * scenario.pages
            input_shape: dict[str, Any] = {
                "format": scenario.format,
                "page_or_slide_count": scenario.pages,
                "page_size_points": {"width": 960, "height": 540},
                "normalizer_render_dpi": 220,
                "expected_render_size_px_per_page": {"width": render_width, "height": render_height},
                "expected_render_pixels_total": expected_pixels,
            }
        else:
            expected_pixels = scenario.width_px * scenario.height_px
            input_shape = {
                "format": "png",
                "page_or_slide_count": 1,
                "canvas_size_px": {"width": scenario.width_px, "height": scenario.height_px},
                "expected_render_pixels_total": expected_pixels,
            }
        if len(durations) == repetitions:
            measurement_status = "MEASURED_REFERENCE_PRODUCT_PATH"
            qualification = "UNQUALIFIED_FULL_PATH_EXTERNAL_MODEL_AND_HOST_STAGES_NOT_MEASURED"
        elif durations:
            measurement_status = "PARTIALLY_MEASURED_REFERENCE_CORE"
            qualification = "UNQUALIFIED_REPETITION_FAILED_ON_TYPED_PRODUCT_PREREQUISITE"
        else:
            measurement_status = "NOT_MEASURED"
            qualification = "NOT_MEASURED_REFERENCE_PATH_PREREQUISITE_UNAVAILABLE"
        results.append({
            "scenario": {
                "scenario_id": scenario.scenario_id,
                "scenario_class": scenario.scenario_class,
                "input_shape": input_shape,
                "input_bytes": source_template.stat().st_size,
            },
            "topology": topology,
            "measurement_status": measurement_status,
            "host_path": {
                "local_cli_reference": "k_slide.cli.main prepare with explicit input paths",
                "opencode_shared_core_reference": "OpenCode host-input reference contract through k_slide.cli.main prepare; plugin lifecycle and model call excluded",
                "cloud_vscode_shared_cli_reference": "Cloud VS Code shared-CLI host-input contract through k_slide.cli.main prepare; external adapter excluded",
            }[topology],
            "shared_engine_path": "prepare_run -> normalize_run -> extract_run; CLI telemetry and product error/state semantics enabled",
            "warmup_policy": {"fresh_product_run_discarded": warmups, "repetitions": repetitions},
            "concurrency": 1,
            "samples": all_samples,
            "observed_latency_ms": {"p50": _percentile(durations, 0.50), "p95": _percentile(durations, 0.95)},
            "throughput_slides_per_second": (scenario.pages * repetitions / (sum(durations) / 1000)) if durations and sum(durations) else None,
            "qualification": qualification,
            "provider_model_invocation": "NOT_MEASURED",
        })
    return results


def _paas_admission_once(root: Path, scenario: Scenario, runtime: RuntimeIdentity, budget: ResourceBudget, run_index: int) -> float:
    scope = AuthorizedScopeContext("ksa36-load-user", "ksa36-load-workspace", "ksa36-load-scope")
    service = ReferencePaaSJobService(root, admission_policy=ScopedAdmissionPolicy.from_resource_budget(budget), scope_context=scope)
    controller = PaaSController(service, scope_context=scope)
    request = PaaSJobRequest(
        f"ksa36-{scenario.scenario_id}-{run_index}",
        str(scope.scope_ref),
        f"ksa36-store-{scenario.scenario_id}-{run_index}",
        runtime,
        total_work_units=scenario.pages,
        scope_context=scope,
    )
    started = time.perf_counter_ns()
    controller.submit(request)
    return (time.perf_counter_ns() - started) / 1_000_000


def _paas_results(output_root: Path, scenario: Scenario, warmups: int, repetitions: int, runtime: RuntimeIdentity, budget: ResourceBudget) -> dict[str, Any]:
    for index in range(warmups):
        with tempfile.TemporaryDirectory(prefix="k-slide-ksa36-paas-warmup-") as temporary:
            _paas_admission_once(Path(temporary), scenario, runtime, budget, index)
    samples: list[float] = []
    batches: list[dict[str, Any]] = []
    for repetition in range(repetitions):
        with tempfile.TemporaryDirectory(prefix="k-slide-ksa36-paas-measure-") as temporary:
            root = Path(temporary)
            scope = AuthorizedScopeContext("ksa36-load-user", "ksa36-load-workspace", "ksa36-load-scope")
            service = ReferencePaaSJobService(root, admission_policy=ScopedAdmissionPolicy.from_resource_budget(budget), scope_context=scope)
            controller = PaaSController(service, scope_context=scope)

            def submit(index: int) -> tuple[float, str]:
                request = PaaSJobRequest(
                    f"ksa36-{scenario.scenario_id}-r{repetition}-j{index}",
                    str(scope.scope_ref),
                    f"ksa36-store-{scenario.scenario_id}-r{repetition}-j{index}",
                    runtime,
                    total_work_units=scenario.pages,
                    scope_context=scope,
                )
                started = time.perf_counter_ns()
                receipt = controller.submit(request)
                return (time.perf_counter_ns() - started) / 1_000_000, receipt.status.value

            batch_started = time.perf_counter_ns()
            with ThreadPoolExecutor(max_workers=4) as pool:
                outcomes = list(pool.map(submit, range(4)))
            batch_ms = (time.perf_counter_ns() - batch_started) / 1_000_000
            samples.extend(item[0] for item in outcomes)
            queue = service.queue(scope_context=scope)
            batches.append({
                "repetition": repetition + 1,
                "requests": len(outcomes),
                "accepted": sum(item[1] in {"ACCEPTED", "IDEMPOTENT"} for item in outcomes),
                "batch_duration_ms": batch_ms,
                "active_runs": sum(entry.active for entry in queue.entries),
                "queued_runs": len(queue.queued),
            })
    return {
        "scenario": {
            "scenario_id": scenario.scenario_id,
            "scenario_class": scenario.scenario_class,
            "input_shape": {
                "format": scenario.format,
                "page_or_slide_count": scenario.pages,
                "authoritative_request_fields": ["total_work_units", "run_ref", "store_ref", "runtime_identity", "authorized_scope"],
                "source_payload": "not_carried_by_the_PaaS_admission_request",
            },
            "input_bytes": None,
            "input_bytes_unavailable_reason": "The authoritative PaaS controller request carries work-unit count and opaque references, not source bytes.",
        },
        "topology": "reference_paas_scoped_admission",
        "measurement_status": "MEASURED_REFERENCE_ADMISSION_ONLY",
        "host_path": "PaaSController.submit -> ReferencePaaSJobService scoped admission",
        "warmup_policy": {"fresh_admission_discarded": warmups, "repetitions": repetitions},
        "concurrency": 4,
        "samples": batches,
        "observed_latency_ms": {"per_submission_p50": _percentile(samples, 0.50), "per_submission_p95": _percentile(samples, 0.95)},
        "throughput_submissions_per_second": (4 * repetitions / (sum(item["batch_duration_ms"] for item in batches) / 1000)) if batches and sum(item["batch_duration_ms"] for item in batches) else None,
        "capacity_outcome": {"configured_active_runs_per_scope": budget.limit("max_concurrent_runs_per_scope"), "configured_queue_depth_per_scope": budget.limit("max_queued_runs_per_scope"), "observed_active_and_queue": batches[-1] if batches else None},
        "excluded_stages": ["company PaaS transport", "source intake and immutable snapshot", "normalization", "evidence extraction", "worker slide engine", "Gemma invocation", "verification", "finalization"],
        "qualification": "UNQUALIFIED_ADMISSION_ONLY_REFERENCE_SERVICE_IS_NOT_COMPANY_PAAS_OR_SLIDE_PROCESSOR",
    }


def run_measurement(root: Path, output: Path, *, warmups: int = 1, repetitions: int = 2) -> dict[str, Any]:
    if not 0 <= warmups <= MAX_WARMUPS or not 1 <= repetitions <= MAX_REPETITIONS:
        raise ValueError("warmups/repetitions are outside the bounded harness range")
    root = root.expanduser().resolve()
    subprocess.run(["git", "merge-base", "--is-ancestor", BASE_SHA, "HEAD"], cwd=root, check=True, capture_output=True)
    if _git(root, "status", "--porcelain"):
        raise ValueError("measurement requires a clean candidate worktree; write reports outside the repository")
    candidate_sha = _git(root, "rev-parse", "HEAD")
    candidate_tree = _git(root, "rev-parse", "HEAD^{tree}")
    changed_paths = _git(root, "diff", "--name-only", f"{BASE_SHA}..{candidate_sha}").splitlines()
    environment = _environment()
    budget = ResourceBudget.reference()
    run_environment = _reference_environment(candidate_sha)
    runtime_identity = RuntimeIdentity(
        "ksa36-local-reference-runtime",
        "NOT_INVOKED",
        "native-extraction-only",
        "core-reference-termbase",
        environment_identity=run_environment,
    )
    results: list[dict[str, Any]] = []
    with tempfile.TemporaryDirectory(prefix="k-slide-ksa36-fixtures-") as fixture_dir:
        fixture_root = Path(fixture_dir)
        for scenario in SCENARIOS:
            source = fixture_root / f"{scenario.scenario_id}.{scenario.format}"
            _write_scenario(source, scenario)
            results.extend(_core_results(fixture_root, scenario, source, warmups, repetitions, run_environment))
            results.append(_paas_results(fixture_root, scenario, warmups, repetitions, runtime_identity, budget))
    measurement_environment_identity = _environment_identity(environment)
    for item in results:
        item["measurement_environment_identity_sha256"] = measurement_environment_identity
        item["kslide_runtime_environment_identity_sha256"] = run_environment.identity_sha256
        if item["topology"] == "reference_paas_scoped_admission":
            item["runtime_identity"] = runtime_identity.as_dict()
        else:
            item["runtime_identity"] = {
                "runtime_ref": "ksa36-local-reference-runtime",
                "model_identity": "NOT_INVOKED",
                "ocr_identity": "native-extraction-only",
                "termbase_identity": "core-reference-termbase",
            }
    report = {
        "schema_version": "1.0",
        "report_type": "k-slide-performance-load-measurement",
        "project": "k-slide",
        "operation": "BUILD",
        "request_id": REQUEST_ID,
        "base_sha": BASE_SHA,
        "branch": _git(root, "branch", "--show-current"),
        "candidate_sha": candidate_sha,
        "candidate_tree": candidate_tree,
        "changed_paths": changed_paths,
        "measurement_environment": environment,
        "measurement_environment_identity_sha256": measurement_environment_identity,
        "kslide_reference_run_environment_identity_sha256": run_environment.identity_sha256,
        "resource_budget": budget.as_dict(),
        "resource_budget_sha256": budget.sha256,
        "production_slo": "NOT_SET_BY_THIS_REPORT",
        "measurement_policy": {
            "warmups": warmups,
            "repetitions": repetitions,
            "maximum_repetitions": MAX_REPETITIONS,
            "core_concurrency": 1,
            "measurement_order": "scenario_then_topology",
            "fixture_policy": "deterministic_synthetic_inputs_deleted_after_run",
            "operational_telemetry_source_content": "NOT_WRITTEN",
        },
        "topology_prerequisites": {
            "local_cli_reference": "REFERENCE_CORE_MEASURED; target model and provider stages NOT_MEASURED",
            "opencode": "SHARED_CORE_MEASURED; actual OpenCode lifecycle, Gemma provider, and provider token usage NOT_MEASURED",
            "cloud_vscode": "SHARED_CLI_CORE_MEASURED; external Cloud VS Code adapter is not present in this repository",
            "company_paas": "NOT_MEASURED; reference scoped admission is measured and reference worker is not a slide engine",
            "pptx_office_converter": "PPTX pipeline is NOT_MEASURED when soffice or libreoffice is absent; no substitute converter is used",
        },
        "scenario_results": results,
        "overall_qualification": "UNQUALIFIED_EXTERNAL_MODEL_CLOUD_HOST_AND_COMPANY_PAAS_STAGES_NOT_MEASURED",
    }
    output = output.expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, sort_keys=True, indent=2) + "\n", encoding="utf-8")
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--warmups", type=int, default=1)
    parser.add_argument("--repetitions", type=int, default=2)
    args = parser.parse_args(argv)
    try:
        report = run_measurement(args.root, args.output, warmups=args.warmups, repetitions=args.repetitions)
    except (KSlideError, OSError, RuntimeError, ValueError, subprocess.CalledProcessError) as exc:
        print(json.dumps({"status": "NOT_MEASURED", "error_type": type(exc).__name__, "recovery": "Install the supported reference dependencies, restore the clean candidate worktree, and rerun with output outside the repository."}, sort_keys=True))
        return 2
    print(json.dumps({"status": "MEASURED_WITH_UNQUALIFIED_EXTERNAL_STAGES", "report": str(args.output.expanduser().resolve()), "report_sha256": _sha256(args.output.expanduser().resolve()), "scenario_results": len(report["scenario_results"]), "candidate_sha": report["candidate_sha"]}, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
