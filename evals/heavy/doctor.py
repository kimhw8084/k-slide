"""Self-test for the heavyweight LibreOffice/Paddle evaluation image."""

from __future__ import annotations

import importlib.util
import json
import importlib.metadata
import os
import platform
import shutil
import subprocess
import tempfile
from pathlib import Path


def _command_version(command: str) -> str | None:
    path = shutil.which(command)
    if not path:
        return None
    try:
        result = subprocess.run([path, "--version"], capture_output=True, text=True, timeout=10, check=False)
    except (OSError, subprocess.TimeoutExpired):
        return None
    from k_slide.certification import parse_libreoffice_version

    return parse_libreoffice_version((result.stdout or "") + "\n" + (result.stderr or "")) if result.returncode == 0 else None


def _package_version(name: str) -> str | None:
    try:
        return importlib.metadata.version(name)
    except importlib.metadata.PackageNotFoundError:
        return None


def _font_check() -> dict[str, object]:
    try:
        from evals.fonts import discover_korean_font

        info = discover_korean_font()
        return {"status": "PASS", "family": info.family, "path": info.path}
    except Exception as exc:
        return {"status": "FAIL", "reason": str(exc)}


def _ocr_asset_identity(*, required: bool = False) -> dict[str, object]:
    """Verify the selected local PaddleX configuration and its model tree."""

    from k_slide.certification import paddle_runtime_configuration, validate_ocr_asset_manifest

    root = Path(os.environ.get("KSLIDE_OCR_ASSET_ROOT", "/opt/k-slide/ocr")).expanduser()
    manifest = root / "manifest.json"
    if not manifest.is_file():
        return {"status": "FAIL" if required else "BLOCKED", "reason": "local OCR asset manifest is unavailable"}
    try:
        identity = validate_ocr_asset_manifest(manifest, asset_root=root, require_model_identity=True)
        configuration = paddle_runtime_configuration(asset_root=root, manifest_path=manifest, require_local=True)
        return {
            "status": "PASS",
            "manifest_sha256": identity["sha256"],
            "file_count": identity["file_count"],
            "paddlex_config": configuration["paddlex_config"],
            "paddlex_config_sha256": configuration["paddlex_config_sha256"],
            "detector_model": configuration.get("detector_model"),
            "recognizer_model": configuration.get("recognizer_model"),
            "offline_assets_required": configuration["offline_assets_required"],
        }
    except Exception as exc:
        return {"status": "FAIL", "reason": str(exc)}


def _libreoffice_roundtrip(*, required: bool = False) -> dict[str, object]:
    if not (_command_version("libreoffice") or _command_version("soffice")):
        return {"status": "FAIL" if required else "BLOCKED", "reason": "LibreOffice/soffice is unavailable."}
    if not all(importlib.util.find_spec(name) for name in ("fitz", "pptx", "PIL")):
        return {"status": "FAIL" if required else "BLOCKED", "reason": "PyMuPDF, python-pptx, and Pillow are required."}
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from k_slide.ingest import prepare_run
        from k_slide.normalization import normalize_run

        with tempfile.TemporaryDirectory(prefix="k-slide-heavy-pptx-") as directory:
            root = Path(directory)
            source = root / "roundtrip.pptx"
            presentation = Presentation()
            blank = presentation.slide_layouts[6]
            for index, (title, body) in enumerate(
                (
                    ("AI Platform", "운영 검토 필요"),
                    ("실적 현황", "3.2조원 / +2.3%p"),
                    ("추진 계획", "검토 후 추진 예정"),
                ),
                start=1,
            ):
                slide = presentation.slides.add_slide(blank)
                title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.6), Inches(11.0), Inches(0.8))
                title_frame = title_box.text_frame
                title_frame.text = title
                title_frame.paragraphs[0].font.size = Pt(28)
                body_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(10.0), Inches(1.0))
                body_box.text_frame.text = f"{index}. {body}"
                body_box.text_frame.paragraphs[0].font.size = Pt(22)
            presentation.save(source)
            run = prepare_run(root / "workspace", explicit_paths=[str(source)])
            normalized = normalize_run(run)
            units = [unit for document in normalized.documents for unit in document.units]
            renders = [run / unit.canonical_render_path for unit in units]
            valid = len(units) == 3 and all(path.is_file() and path.stat().st_size > 0 for path in renders)
            return {"status": "PASS" if valid else "FAIL", "slide_count": len(units), "render_count": len(renders), "render_dimensions": [[unit.width_px, unit.height_px] for unit in units]}
    except Exception as exc:
        return {"status": "FAIL", "reason": str(exc)}


def _paddle_ocr_roundtrip(*, required: bool = False) -> dict[str, object]:
    missing = [name for name, module in (("PaddleOCR", "paddleocr"), ("PaddlePaddle", "paddle"), ("Pillow", "PIL")) if importlib.util.find_spec(module) is None]
    if missing:
        return {"status": "FAIL" if required else "BLOCKED", "reason": f"Unavailable: {', '.join(missing)}."}
    try:
        from PIL import Image, ImageDraw
        from evals.fonts import discover_korean_font, load_font
        from k_slide.ocr.policy import OCRProviderPolicy, create_ocr_provider

        font_info = discover_korean_font()
        with tempfile.TemporaryDirectory(prefix="k-slide-heavy-ocr-") as directory:
            image_path = Path(directory) / "korean-smoke.png"
            image = Image.new("RGB", (1600, 900), "white")
            draw = ImageDraw.Draw(image)
            font = load_font(42, font_info)
            for index, text in enumerate(("운영 검토 필요", "3.2조원", "+2.3%p", "검토 후 추진 예정")):
                draw.text((80, 100 + index * 150), text, fill="black", font=font)
            image.save(image_path)
            selection = create_ocr_provider(OCRProviderPolicy.PADDLE)
            result = selection.provider.extract(image_path)
            texts = " ".join(region.text for region in result.regions)
            expected = ("검토", "3.2", "%p")
            found = sum(term in texts for term in expected)
            boxes_valid = all(region.bbox_px[2] > region.bbox_px[0] and region.bbox_px[3] > region.bbox_px[1] for region in result.regions)
            return {"status": "PASS" if result.regions and found >= 1 and boxes_valid else "FAIL", "region_count": len(result.regions), "expected_key_recall": found / len(expected), "boxes_valid": boxes_valid, "provider": result.provider, "version": result.provider_version, "requested_policy": selection.requested, "effective_provider": selection.effective}
    except Exception as exc:
        return {"status": "FAIL", "reason": str(exc)}


def main(*, required: bool = False, networkless: bool = False) -> int:
    def available(module: str) -> str:
        return "PASS" if importlib.util.find_spec(module) else ("FAIL" if required else "BLOCKED")

    checks: dict[str, object] = {
        "python": {"status": "PASS"},
        "libreoffice": {"status": "PASS" if _command_version("libreoffice") or _command_version("soffice") else ("FAIL" if required else "BLOCKED")},
        "pillow": {"status": available("PIL")},
        "pymupdf": {"status": available("fitz")},
        "python_pptx": {"status": available("pptx")},
        "paddleocr": {"status": available("paddleocr")},
        "paddlepaddle": {"status": available("paddle")},
        "korean_font": _font_check(),
    }
    checks["runtime_provenance"] = {
        "python_version": platform.python_version(),
        "paddle_version": _package_version("paddlepaddle"),
        "paddleocr_version": _package_version("paddleocr"),
        "libreoffice_version": _command_version("libreoffice") or _command_version("soffice"),
    }
    checks["ocr_asset_identity"] = _ocr_asset_identity(required=required)
    # A small real OCR load check is intentionally separate from package import.
    if checks["paddleocr"].get("status") == "PASS" and checks["pillow"].get("status") == "PASS":
        try:
            from k_slide.ocr.policy import OCRProviderPolicy, create_ocr_provider

            selection = create_ocr_provider(OCRProviderPolicy.PADDLE)
            checks["paddle_load"] = {"status": "PASS", "provider": selection.effective, "version": selection.version, "requested_policy": selection.requested}
        except Exception as exc:
            checks["paddle_load"] = {"status": "FAIL", "reason": str(exc)}
    else:
        missing = [name for name, module in (("PaddleOCR", "paddleocr"), ("Pillow", "PIL")) if importlib.util.find_spec(module) is None]
        checks["paddle_load"] = {"status": "FAIL" if required else "BLOCKED", "reason": f"Unavailable: {', '.join(missing)}."}
    checks["libreoffice_roundtrip"] = _libreoffice_roundtrip(required=required)
    checks["paddle_ocr_roundtrip"] = _paddle_ocr_roundtrip(required=required)
    checks["network"] = {"network_required": not networkless, "networkless_asserted": networkless}
    print(json.dumps(checks, ensure_ascii=False, indent=2))
    if required:
        required_checks = ("libreoffice", "paddleocr", "paddlepaddle", "korean_font", "ocr_asset_identity", "paddle_load", "libreoffice_roundtrip", "paddle_ocr_roundtrip")
        return 0 if all(checks.get(name, {}).get("status") == "PASS" for name in required_checks) else 1
    status_values = [value.get("status") for value in checks.values() if isinstance(value, dict) and "status" in value]
    return 0 if all(status in {"PASS", "BLOCKED"} for status in status_values) else 1


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="K-Slide heavy runtime doctor")
    parser.add_argument("--required", action="store_true", help="fail when heavy dependencies or round trips are unavailable")
    parser.add_argument("--networkless", action="store_true", help="record that this check is expected to run without network; the container boundary enforces this")
    args = parser.parse_args()
    raise SystemExit(main(required=args.required, networkless=args.networkless))
