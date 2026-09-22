"""Deterministic image/PDF/PPTX normalization with explicit capability failures."""

from __future__ import annotations

import importlib.util
import os
import shutil
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

from .documents import DocumentUnit, NormalizationResult, NormalizedDocument
from .environment import RunEnvironmentIdentity
from .errors import ErrorCode, KSlideError
from .evidence_ir import stable_revision
from .io import atomic_write_json, atomic_write_text, read_json
from .locking import run_lock
from .queue import WorkUnit, WorkUnitStatus, WorkQueue, load_queue, save_queue
from .security import sha256_file
from .storage import StorageArtifact, StorageLayout, storage_path
from .state import RunPhase, load_state, save_state
from .redaction import safe_diagnostic_text_or_placeholder
def _terminate_subprocess_group(process: subprocess.Popen[str], grace_seconds: float = 5.0) -> None:
    """Keep document conversion cleanup local to the core runtime."""

    import signal

    try:
        os.killpg(process.pid, signal.SIGTERM)
    except (OSError, ProcessLookupError):
        try:
            process.terminate()
        except OSError:
            return
    try:
        process.wait(timeout=grace_seconds)
    except subprocess.TimeoutExpired:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except (OSError, ProcessLookupError):
            try:
                process.kill()
            except OSError:
                pass
        try:
            process.wait(timeout=grace_seconds)
        except subprocess.TimeoutExpired:
            pass

RENDER_DPI = 220.0
MAX_IMAGE_PIXELS = 120_000_000
MAX_IMAGE_WIDTH = 20_000
MAX_IMAGE_HEIGHT = 20_000
MAX_DOCUMENTS_PER_RUN = 32
MAX_UNITS_PER_RUN = 500
MAX_PAGES_PER_PDF = 200
MAX_SLIDES_PER_PPTX = 200
MAX_TOTAL_RENDER_PIXELS = 500_000_000
MAX_NORMALIZED_BYTES = 2_000_000_000


def _json_safe(value: Any) -> Any:
    """Keep PDF layout metadata serializable without persisting image bytes."""

    if isinstance(value, bytes):
        return {"byte_length": len(value), "omitted": True}
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items() if key not in {"image", "mask"}}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    return value


def _snapshot_inputs(run_dir: Path) -> list[tuple[str, Path, dict[str, Any]]]:
    manifest = read_json(storage_path(run_dir, StorageArtifact.RUN_MANIFEST, "RUN_MANIFEST.json"))
    values = manifest.get("inputs", []) if isinstance(manifest, dict) else []
    results: list[tuple[str, Path, dict[str, Any]]] = []
    for index, item in enumerate(values, start=1):
        extension = str(item.get("extension", "")).lower()
        if extension not in {".png", ".jpg", ".jpeg", ".webp", ".pdf", ".pptx"}:
            raise KSlideError(ErrorCode.INPUT_CORRUPT, "Immutable input snapshot metadata is invalid.")
        source = storage_path(run_dir, StorageArtifact.SOURCE_SNAPSHOT, f"inputs/source-{index:03d}{extension}")
        if not source.is_file() or source.is_symlink():
            raise KSlideError(ErrorCode.INPUT_NOT_FOUND, "Immutable input snapshot is missing.", {"input_id": f"source-{index:03d}"})
        expected_sha256 = item.get("sha256")
        if not isinstance(expected_sha256, str) or sha256_file(source) != expected_sha256:
            raise KSlideError(ErrorCode.INPUT_CORRUPT, "Immutable input snapshot failed hash verification.")
        results.append((f"source-{index:03d}", source, {**item, "extension": extension}))
    return results


def _pptx_slide_count(source: Path) -> int:
    """Preflight slide cardinality before invoking an Office converter."""

    try:
        with zipfile.ZipFile(source) as archive:
            data = archive.read("ppt/presentation.xml")
        root = ET.fromstring(data)
    except (ET.ParseError, OSError, ValueError, zipfile.BadZipFile, KeyError) as exc:
        raise KSlideError(ErrorCode.NORMALIZATION_FAILED, "PPTX package structure could not be inspected safely.") from exc
    return sum(1 for item in root.iter() if item.tag.rsplit("}", 1)[-1].casefold() == "sldid")


def _normalize_image(run_dir: Path, input_id: str, source: Path, index: int, document_id: str) -> list[DocumentUnit]:
    if importlib.util.find_spec("PIL") is None:
        raise KSlideError(ErrorCode.IMAGE_DECODE_FAILED, "Pillow is required to decode and normalize image inputs.", {"input": source.name, "extra": "pip install k-slide"})
    from PIL import Image, ImageOps, UnidentifiedImageError
    DecompressionBombError = getattr(Image, "DecompressionBombError", OSError)

    try:
        with Image.open(source) as image:
            image = ImageOps.exif_transpose(image)
            width, height = image.size
            if width < 1 or height < 1 or width > MAX_IMAGE_WIDTH or height > MAX_IMAGE_HEIGHT or width * height > MAX_IMAGE_PIXELS:
                raise KSlideError(ErrorCode.INPUT_TOO_LARGE, "Decoded image dimensions exceed K-Slide safety limits.", {"width": width, "height": height, "max_pixels": MAX_IMAGE_PIXELS})
            if image.mode not in {"RGB", "RGBA"}:
                image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
            work_unit_id = f"{document_id}-image-0001"
            render = storage_path(run_dir, StorageArtifact.NORMALIZED_RENDER, f"normalized/{work_unit_id}.png", create_parent=True)
            image.save(render, format="PNG", optimize=True)
            render.chmod(0o600)
    except KSlideError:
        raise
    except (OSError, UnidentifiedImageError, DecompressionBombError, ValueError) as exc:
        raise KSlideError(ErrorCode.IMAGE_DECODE_FAILED, "Image could not be decoded safely.", {"input": source.name, "reason": type(exc).__name__}) from exc
    native_path = storage_path(run_dir, StorageArtifact.NATIVE_EXTRACTION, f"native/{work_unit_id}.json", create_parent=True)
    atomic_write_json(native_path, {"provider": "image_decoder", "regions": []}, mode=0o600)
    durable_root = Path(run_dir).resolve()
    return [DocumentUnit(work_unit_id, document_id, input_id, index - 1, "image", width, height, str(render.resolve().relative_to(durable_root)), sha256_file(render), str(native_path.relative_to(durable_root)), ())]


def _pdf_units(run_dir: Path, input_id: str, source: Path, document_id: str, source_index: int) -> list[DocumentUnit]:
    if importlib.util.find_spec("fitz") is None:
        raise KSlideError(ErrorCode.PDF_RENDER_UNAVAILABLE, "PyMuPDF is required to normalize PDF pages.", {"input": source.name, "extra": "pip install k-slide[pdf]"})
    import fitz

    try:
        document = fitz.open(source)
    except Exception as exc:
        raise KSlideError(ErrorCode.NORMALIZATION_FAILED, "PDF could not be opened.", {"input": source.name, "reason": type(exc).__name__}) from exc
    if document.is_encrypted:
        document.close()
        raise KSlideError(ErrorCode.PDF_ENCRYPTED, "Password-protected PDFs are not accepted by the local normalizer.", {"input": source.name})
    if document.page_count < 1:
        document.close()
        raise KSlideError(ErrorCode.NORMALIZATION_FAILED, "PDF contains no pages.", {"input": source.name})
    if document.page_count > MAX_PAGES_PER_PDF:
        page_count = document.page_count
        document.close()
        raise KSlideError(ErrorCode.RESOURCE_LIMIT, "PDF page count exceeds the configured safety limit.", {"input": source.name, "pages": page_count, "max_pages": MAX_PAGES_PER_PDF})
    units: list[DocumentUnit] = []
    try:
        for page_index in range(document.page_count):
            page = document.load_page(page_index)
            rect = page.rect
            matrix = fitz.Matrix(RENDER_DPI / 72.0, RENDER_DPI / 72.0)
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            work_unit_id = f"{document_id}-page-{page_index + 1:04d}"
            render = storage_path(run_dir, StorageArtifact.NORMALIZED_RENDER, f"normalized/{work_unit_id}.png", create_parent=True)
            pixmap.save(str(render))
            render.chmod(0o600)
            native = page.get_text("dict")
            native_path = storage_path(run_dir, StorageArtifact.NATIVE_EXTRACTION, f"native/{work_unit_id}.json", create_parent=True)
            safe_blocks = _json_safe(native.get("blocks", []))
            atomic_write_json(native_path, {"provider": "pymupdf", "page_index": page_index, "blocks": safe_blocks}, mode=0o600)
            durable_root = Path(run_dir).resolve()
            units.append(DocumentUnit(work_unit_id, document_id, input_id, page_index, "page", pixmap.width, pixmap.height, str(render.resolve().relative_to(durable_root)), sha256_file(render), str(native_path.relative_to(durable_root)), tuple(safe_blocks), RENDER_DPI))
    finally:
        document.close()
    return units


def _shape_bbox(shape: Any, width_emu: int, height_emu: int, width_px: int, height_px: int) -> tuple[int, int, int, int]:
    """Map PPTX EMU geometry into the canonical render pixel space."""

    x = round(int(getattr(shape, "left", 0)) / max(width_emu, 1) * width_px)
    y = round(int(getattr(shape, "top", 0)) / max(height_emu, 1) * height_px)
    w = round(int(getattr(shape, "width", width_emu)) / max(width_emu, 1) * width_px)
    h = round(int(getattr(shape, "height", height_emu)) / max(height_emu, 1) * height_px)
    return x, y, max(x + w, x + 1), max(y + h, y + 1)


def _iter_shapes(shapes: Any, prefix: str = "") -> list[tuple[str, Any]]:
    """Traverse grouped PPTX shapes with stable source numbering."""

    result: list[tuple[str, Any]] = []
    for index, shape in enumerate(shapes, start=1):
        source_index = f"{prefix}{index:03d}"
        result.append((source_index, shape))
        children = getattr(shape, "shapes", None)
        if children is not None:
            result.extend(_iter_shapes(children, f"{source_index}-"))
    return result


def _enum_name(value: Any, default: str = "unknown") -> str:
    name = getattr(value, "name", None)
    if not name:
        name = str(value)
    name = str(name).split("(", 1)[0].strip().lower()
    return name or default


def _safe_shape_text(shape: Any) -> str | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    try:
        return "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs)
    except (AttributeError, TypeError, ValueError):
        return None


def _axis_text(axis: Any) -> str | None:
    try:
        return axis.axis_title.text_frame.text if axis.has_title else None
    except (AttributeError, ValueError):
        return None


def _chart_metadata(shape: Any) -> dict[str, Any] | None:
    if not getattr(shape, "has_chart", False):
        return None
    chart = shape.chart
    value: dict[str, Any] = {"structure_version": "1.0", "chart_type": _enum_name(getattr(chart, "chart_type", "unknown"))}
    try:
        value["title"] = chart.chart_title.text_frame.text if chart.has_title else None
    except (AttributeError, ValueError):
        value["title"] = None
    series_values: list[dict[str, Any]] = []
    for series_index, series in enumerate(getattr(chart, "series", [])):
        name = getattr(series, "name", "")
        item: dict[str, Any] = {"series_index": series_index, "name": str(name) if name is not None else ""}
        try:
            raw_values = list(series.values)
            points = []
            for index, number in enumerate(raw_values):
                if number is None or number == "":
                    points.append({"point_index": index, "value": None, "is_blank": True})
                else:
                    points.append({"point_index": index, "value": float(number), "is_blank": False})
            item["points"] = points
            # Keep the legacy field for readers written before KSA-24.
            item["values"] = [point["value"] for point in points]
        except (AttributeError, TypeError, ValueError):
            item["points"] = []
            item["values"] = []
        series_values.append(item)
    value["series"] = series_values
    value["series_order"] = [item["name"] for item in series_values]
    try:
        value["categories"] = [None if category is None else str(category) for category in chart.plots[0].categories]
    except (AttributeError, IndexError, TypeError):
        value["categories"] = []
    try:
        category_axis = chart.category_axis
        value_axis = chart.value_axis
        value["axis_labels"] = {"category": _axis_text(category_axis), "value": _axis_text(value_axis)}
        value["unit_labels"] = {"category": None, "value": getattr(value_axis.tick_labels, "number_format", None)}
    except (AttributeError, ValueError):
        value["axis_labels"] = {"category": None, "value": None}
        value["unit_labels"] = {"category": None, "value": None}
    return value


def _table_header_flags(table: Any) -> tuple[bool | None, bool | None]:
    """Read explicit PowerPoint first-row/first-column table semantics."""

    try:
        properties = table._tbl.tblPr
        def flag(name: str) -> bool | None:
            value = properties.get(name)
            if value is None:
                return None
            return str(value).lower() in {"1", "true", "on"}
        return flag("firstRow"), flag("firstCol")
    except (AttributeError, TypeError, ValueError):
        return None, None


def _connector_endpoints(shape: Any, source_ids_by_shape_id: dict[int, str]) -> dict[str, Any] | None:
    try:
        element = shape._element
        properties = next((item for item in element.iter() if item.tag.rsplit("}", 1)[-1] == "cNvCxnSpPr"), None)
        if properties is None:
            return None
        values: dict[str, int] = {}
        for child in properties:
            local = child.tag.rsplit("}", 1)[-1]
            if local in {"stCxn", "endCxn"} and child.get("id") is not None:
                values[local] = int(child.get("id"))
        if "stCxn" not in values or "endCxn" not in values:
            return None
        start = source_ids_by_shape_id.get(values["stCxn"])
        end = source_ids_by_shape_id.get(values["endCxn"])
        if not start or not end:
            return None
        return {"from_element_id": start, "to_element_id": end, "start_shape_id": values["stCxn"], "end_shape_id": values["endCxn"]}
    except (AttributeError, TypeError, ValueError):
        return None


def _pptx_native(source: Path, run_dir: Path, input_id: str, document_id: str, rendered_pages: list[Path]) -> list[DocumentUnit]:
    if importlib.util.find_spec("pptx") is None:
        raise KSlideError(ErrorCode.NORMALIZATION_FAILED, "python-pptx is required to extract PPTX structure.", {"input": source.name, "extra": "pip install k-slide[pptx]"})
    from pptx import Presentation

    try:
        presentation = Presentation(str(source))
    except Exception as exc:
        raise KSlideError(ErrorCode.NORMALIZATION_FAILED, "PPTX could not be parsed.", {"input": source.name, "reason": type(exc).__name__}) from exc
    slide_width_emu = int(presentation.slide_width)
    slide_height_emu = int(presentation.slide_height)
    units: list[DocumentUnit] = []
    if len(rendered_pages) != len(presentation.slides):
        raise KSlideError(ErrorCode.RENDER_COUNT_MISMATCH, "PPTX render page count does not match slide count.", {"slides": len(presentation.slides), "renders": len(rendered_pages)})
    if len(presentation.slides) > MAX_SLIDES_PER_PPTX:
        raise KSlideError(ErrorCode.RESOURCE_LIMIT, "PPTX slide count exceeds the configured safety limit.", {"input": source.name, "slides": len(presentation.slides), "max_slides": MAX_SLIDES_PER_PPTX})
    for slide_index, slide in enumerate(presentation.slides):
        work_unit_id = f"{document_id}-slide-{slide_index + 1:04d}"
        render = rendered_pages[slide_index]
        try:
            from PIL import Image

            # The rendered PNG dimensions are authoritative for crop geometry.
            with Image.open(render) as rendered_image:
                render_width, render_height = rendered_image.size
        except (ImportError, OSError) as exc:
            raise KSlideError(ErrorCode.PPTX_RENDER_UNAVAILABLE, "Unable to inspect the canonical PPTX render dimensions.", {"input": source.name}) from exc
        objects: list[dict[str, Any]] = []
        shape_records = _iter_shapes(slide.shapes)
        source_ids_by_shape_id = {
            int(getattr(shape, "shape_id")): f"{work_unit_id}-shape-{shape_index}"
            for shape_index, shape in shape_records
            if getattr(shape, "shape_id", None) is not None
        }
        for shape_index, shape in shape_records:
            source_id = f"{work_unit_id}-shape-{shape_index}"
            shape_type = _enum_name(getattr(shape, "shape_type", "unknown"))
            item: dict[str, Any] = {
                "source_id": source_id,
                "shape_id": getattr(shape, "shape_id", None),
                "shape_type": shape_type,
                "bbox_px": list(_shape_bbox(shape, slide_width_emu, slide_height_emu, render_width, render_height)),
            }
            text = _safe_shape_text(shape)
            if text is not None:
                item["text"] = text
            if getattr(shape, "has_table", False):
                table = shape.table
                first_row, first_column = _table_header_flags(table)
                cells = []
                for row_index in range(len(table.rows)):
                    for column_index in range(len(table.columns)):
                        cell = table.cell(row_index, column_index)
                        is_origin = bool(getattr(cell, "is_merge_origin", False))
                        is_spanned = bool(getattr(cell, "is_spanned", False))
                        is_blank = not is_spanned and not bool(cell.text)
                        state = "merge_continuation" if is_spanned else ("merge_origin" if is_origin else ("blank" if is_blank else "nonblank"))
                        is_header = None
                        if first_row is not None and row_index == 0:
                            is_header = first_row
                        if first_column is not None and column_index == 0:
                            is_header = bool(is_header or first_column)
                        rowspan = int(getattr(cell, "span_height", 1) or 1) if not is_spanned else 1
                        colspan = int(getattr(cell, "span_width", 1) or 1) if not is_spanned else 1
                        cells.append({
                            "cell_id": f"{work_unit_id}-table-{shape_index}-r{row_index + 1:02d}-c{column_index + 1:02d}",
                            "row": row_index,
                            "column": column_index,
                            "text": None if is_spanned else cell.text,
                            "rowspan": rowspan,
                            "colspan": colspan,
                            "is_merge_origin": is_origin,
                            "is_spanned": is_spanned,
                            "is_blank": is_blank,
                            "is_header": is_header,
                        })
                headers = [cell["text"] for cell in cells if cell["is_header"] is True and cell["text"]]
                item["table"] = {
                    "row_count": len(table.rows),
                    "column_count": len(table.columns),
                    "cells": cells,
                    "first_row_header": first_row,
                    "first_column_header": first_column,
                    "header_rows": [0] if first_row is True else [],
                    "header_columns": [0] if first_column is True else [],
                    "headers": headers,
                }
            chart = _chart_metadata(shape)
            if chart is not None:
                item["chart"] = chart
            if shape_type in {"connector", "line"}:
                connector = _connector_endpoints(shape, source_ids_by_shape_id)
                if connector is not None:
                    item["connector"] = connector
            objects.append(item)
        native_path = storage_path(run_dir, StorageArtifact.NATIVE_EXTRACTION, f"native/{work_unit_id}.json", create_parent=True)
        atomic_write_json(native_path, {"provider": "python-pptx", "slide_index": slide_index, "slide_width_px": render_width, "slide_height_px": render_height, "objects": objects}, mode=0o600)
        durable_root = Path(run_dir).resolve()
        units.append(DocumentUnit(work_unit_id, document_id, input_id, slide_index, "slide", render_width, render_height, str(render.resolve().relative_to(durable_root)), sha256_file(render), str(native_path.relative_to(durable_root)), tuple(objects)))
    return units


def _render_pptx(source: Path, run_dir: Path, document_id: str) -> list[Path]:
    binary = shutil.which("libreoffice") or shutil.which("soffice")
    if not binary:
        raise KSlideError(ErrorCode.PPTX_RENDER_UNAVAILABLE, "LibreOffice/soffice is required to render PPTX slides.", {"input": source.name})
    scratch = StorageLayout.for_workspace(run_dir)
    staging_root = scratch.ensure_directory(StorageArtifact.CONVERSION_STAGING, f"office/{document_id}")
    with tempfile.TemporaryDirectory(prefix="k-slide-office-", dir=staging_root) as temporary:
        output = Path(temporary) / "out"
        profile = Path(temporary) / "profile"
        output.mkdir()
        profile.mkdir()
        command = [binary, "--headless", f"-env:UserInstallation=file://{profile}", "--convert-to", "pdf", "--outdir", str(output), str(source)]
        try:
            process = subprocess.Popen(command, cwd=temporary, env={"PATH": os.environ.get("PATH", ""), "HOME": temporary, "LANG": "C.UTF-8"}, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, start_new_session=True)
        except OSError as exc:
                raise KSlideError(ErrorCode.PPTX_RENDER_UNAVAILABLE, "PPTX converter could not be started.", {"input": source.name, "reason": type(exc).__name__}) from exc
        try:
            stdout, stderr = process.communicate(timeout=120)
        except subprocess.TimeoutExpired as exc:
            _terminate_subprocess_group(process)
            raise KSlideError(ErrorCode.PPTX_CONVERSION_TIMEOUT, "PPTX rendering timed out.", {"input": source.name}) from exc
        if process.returncode != 0:
            raise KSlideError(ErrorCode.PPTX_RENDER_UNAVAILABLE, "PPTX could not be rendered by the available Office converter.", {"input": source.name, "reason": "converter_exit"})
        pdf = output / f"{source.stem}.pdf"
        if not pdf.is_file():
            raise KSlideError(ErrorCode.PPTX_RENDER_UNAVAILABLE, "PPTX converter produced no PDF output.", {"input": source.name})
        if importlib.util.find_spec("fitz") is None:
            raise KSlideError(ErrorCode.PPTX_RENDER_UNAVAILABLE, "PyMuPDF is required to render converted PPTX pages.", {"input": source.name})
        import fitz
        document = fitz.open(pdf)
        try:
            renders: list[Path] = []
            for index, page in enumerate(document):
                render = storage_path(run_dir, StorageArtifact.NORMALIZED_RENDER, f"normalized/{document_id}-slide-{index + 1:04d}.png", create_parent=True)
                page.get_pixmap(matrix=fitz.Matrix(RENDER_DPI / 72, RENDER_DPI / 72), alpha=False).save(str(render))
                render.chmod(0o600)
                renders.append(render)
            return renders
        finally:
            document.close()


def normalize_run(run_dir: Path, *, environment_identity: RunEnvironmentIdentity | None = None) -> NormalizationResult:
    """Normalize immutable run inputs and replace the initial queue with document units."""

    from .execution import ensure_workspace_environment_compatible

    ensure_workspace_environment_compatible(run_dir, environment_identity=environment_identity)
    with run_lock(run_dir):
        state = load_state(run_dir)
        if state.phase not in {RunPhase.INPUT_VALIDATED, RunPhase.FAILED_NORMALIZATION}:
            if state.phase in {RunPhase.NORMALIZED, RunPhase.EXTRACTED, RunPhase.TRANSLATING, RunPhase.TRANSLATED, RunPhase.VERIFYING, RunPhase.VERIFIED, RunPhase.COMPLETE}:
                return _load_normalization_result(run_dir)
            raise KSlideError(ErrorCode.INVALID_TRANSITION, "Run is not ready for normalization.", {"phase": state.phase.value})
        state.transition(RunPhase.NORMALIZING, next_action="Normalize immutable document inputs")
        save_state(run_dir, state)
        documents: list[NormalizedDocument] = []
        warnings: list[str] = []
        try:
            snapshots = _snapshot_inputs(run_dir)
            if len(snapshots) > MAX_DOCUMENTS_PER_RUN:
                raise KSlideError(ErrorCode.RESOURCE_LIMIT, "Input count exceeds the configured safety limit.", {"documents": len(snapshots), "max_documents": MAX_DOCUMENTS_PER_RUN})
            for input_id, source, metadata in snapshots:
                document_id = f"doc-{input_id.split('-')[-1]}"
                extension = str(metadata.get("extension", source.suffix)).lower()
                if extension in {".png", ".jpg", ".jpeg", ".webp"}:
                    units = _normalize_image(run_dir, input_id, source, int(input_id.split("-")[-1]), document_id)
                elif extension == ".pdf":
                    units = _pdf_units(run_dir, input_id, source, document_id, int(input_id.split("-")[-1]))
                elif extension == ".pptx":
                    slide_count = _pptx_slide_count(source)
                    if slide_count > MAX_SLIDES_PER_PPTX:
                        raise KSlideError(ErrorCode.RESOURCE_LIMIT, "PPTX slide count exceeds the configured safety limit.", {"max_slides": MAX_SLIDES_PER_PPTX})
                    rendered = _render_pptx(source, run_dir, document_id)
                    units = _pptx_native(source, run_dir, input_id, document_id, rendered)
                else:
                    raise KSlideError(ErrorCode.INPUT_UNSUPPORTED, "Unsupported normalized input type.", {"extension": extension})
                documents.append(NormalizedDocument(document_id, input_id, str(metadata.get("kind", "unknown")), str(metadata.get("sha256", "")), tuple(units)))
            units = [unit for document in documents for unit in document.units]
            document_ids = [document.document_id for document in documents]
            if len(document_ids) != len(set(document_ids)):
                raise KSlideError(ErrorCode.DUPLICATE_DOCUMENT_ID, "Normalization generated duplicate document IDs.", {"document_ids": document_ids})
            if len(units) > MAX_UNITS_PER_RUN:
                raise KSlideError(ErrorCode.RESOURCE_LIMIT, "Page/slide count exceeds the configured safety limit.", {"units": len(units), "max_units": MAX_UNITS_PER_RUN})
            total_pixels = sum(unit.width_px * unit.height_px for unit in units)
            if total_pixels > MAX_TOTAL_RENDER_PIXELS:
                raise KSlideError(ErrorCode.RESOURCE_LIMIT, "Total normalized render pixels exceed the configured safety limit.", {"pixels": total_pixels, "max_pixels": MAX_TOTAL_RENDER_PIXELS})
            normalized_bytes = sum(storage_path(run_dir, StorageArtifact.NORMALIZED_RENDER, unit.canonical_render_path).stat().st_size for unit in units if storage_path(run_dir, StorageArtifact.NORMALIZED_RENDER, unit.canonical_render_path).is_file())
            if normalized_bytes > MAX_NORMALIZED_BYTES:
                raise KSlideError(ErrorCode.RESOURCE_LIMIT, "Normalized render bytes exceed the configured safety limit.", {"bytes": normalized_bytes, "max_bytes": MAX_NORMALIZED_BYTES})
            result = NormalizationResult(tuple(documents), tuple(warnings), {"render_dpi": RENDER_DPI, "max_image_pixels": MAX_IMAGE_PIXELS})
            atomic_write_json(storage_path(run_dir, StorageArtifact.NORMALIZATION_MANIFEST, "normalized/DOCUMENT_MANIFEST.json", create_parent=True), result.as_dict(), mode=0o600)
            now = state.updated_at
            queue = WorkQueue(run_id=state.run_id, work_units=[WorkUnit(unit.work_unit_id, unit.document_id, unit.input_id, unit.source_index, unit.kind, WorkUnitStatus.NORMALIZED, created_at=now, updated_at=now) for unit in units])
            save_queue(run_dir, queue)
            state.input_count = len(units)
            state.transition(RunPhase.NORMALIZED, next_action="Extract native evidence and generate regions")
            save_state(run_dir, state)
            return result
        except KSlideError as exc:
            state = load_state(run_dir)
            state.transition(RunPhase.FAILED_NORMALIZATION, next_action="Fix the normalization capability or source file and retry", error_code=exc.code.value, error_message=exc.message)
            save_state(run_dir, state)
            atomic_write_json(storage_path(run_dir, StorageArtifact.NORMALIZATION_ERROR, "normalized/NORMALIZATION_ERROR.json", create_parent=True), exc.as_dict(), mode=0o600)
            atomic_write_text(storage_path(run_dir, StorageArtifact.FAILURE_MARKER, "RUN_FAILED.md", create_parent=True), f"# FAILED\n\n{safe_diagnostic_text_or_placeholder(exc.message)}\n\nError code: `{exc.code.value}`\n")
            raise
        except (ImportError, OSError, RuntimeError, ValueError) as exc:
            error = KSlideError(ErrorCode.NORMALIZATION_FAILED, "Document normalization failed safely.", {"reason": type(exc).__name__})
            state = load_state(run_dir)
            state.transition(RunPhase.FAILED_NORMALIZATION, next_action="Fix the normalization capability or source file and retry", error_code=error.code.value, error_message=error.message)
            save_state(run_dir, state)
            atomic_write_json(storage_path(run_dir, StorageArtifact.NORMALIZATION_ERROR, "normalized/NORMALIZATION_ERROR.json", create_parent=True), error.as_dict(), mode=0o600)
            atomic_write_text(storage_path(run_dir, StorageArtifact.FAILURE_MARKER, "RUN_FAILED.md", create_parent=True), f"# FAILED\n\n{safe_diagnostic_text_or_placeholder(error.message)}\n\nError code: `{error.code.value}`\n")
            raise error from exc


def _load_normalization_result(run_dir: Path) -> NormalizationResult:
    value = read_json(storage_path(run_dir, StorageArtifact.NORMALIZATION_MANIFEST, "normalized/DOCUMENT_MANIFEST.json"))
    documents: list[NormalizedDocument] = []
    for item in value.get("documents", []):
        units = tuple(DocumentUnit(**{**unit, "native_evidence": tuple(unit.get("native_evidence", [])), "warnings": tuple(unit.get("warnings", []))}) for unit in item.get("units", []))
        documents.append(NormalizedDocument(item["document_id"], item["input_id"], item["kind"], item["source_sha256"], units))
    return NormalizationResult(tuple(documents), tuple(value.get("warnings", [])), dict(value.get("capabilities", {})))
